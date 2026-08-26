"""add_reverse_dcf_sheet.py - RETROFIT the standard 'Reverse DCF' sheet onto an
already-generated DCF workbook.

The sheet is no longer built here. templates/reverse_dcf_sheet.py owns the whole
implementation and dcf_comps_template writes it as the standard 4th sheet of 8
on every run, so a workbook generated after that change already has it. This
script exists only for workbooks generated BEFORE it — it loads the file, reads
the layout back out of the 'Adjustments Log' Pipeline Metadata block, and calls
the same builder. There is one implementation, not two.

Usage:
    python scripts/add_reverse_dcf_sheet.py models/5726_DCF_Model_20260826.xlsx \
        [--op0 5524 --op0-label FY2026/3 --peak-op 10088 --peak-label FY2025/3 \
         --peak-opm 0.194 --benchmark-ticker 5727 --deal-note "..."]

Every input is optional: whatever is omitted is derived the same way a fresh run
would derive it, from the workbook's own historical operating profit rows.

openpyxl drops the cached formula values of every sheet on save, so run
scripts/recalc_excel_com.py on the workbook afterwards (and re-validate).
"""
import argparse
import os
import sys

import openpyxl

sys.path.insert(0, os.path.abspath(
    os.path.join(os.path.dirname(os.path.abspath(__file__)), "..")))

from templates.dcf_comps_template import (  # noqa: E402
    R_REVENUE, R_DA, R_EV_PGM, col_letter,
)
from templates.reverse_dcf_sheet import (  # noqa: E402
    SHEET_NAME, resolve_params, build_reverse_dcf_sheet, place_after, selftest_bn,
)


def _read_metadata(wb):
    """Read the Pipeline Metadata key/value block from 'Adjustments Log'."""
    meta = {}
    if "Adjustments Log" not in wb.sheetnames:
        return meta
    ws = wb["Adjustments Log"]
    seen = False
    for r in range(1, ws.max_row + 1):
        b = ws.cell(row=r, column=2).value
        if isinstance(b, str) and b.startswith("Pipeline Metadata"):
            seen = True
            continue
        if seen and b not in (None, "", "key"):
            meta[str(b)] = ws.cell(row=r, column=3).value
    return meta


def _projection_last_col(ws3):
    """Last column of the projection block, found from the revenue row itself."""
    last = 3
    for col in range(3, 20):
        if ws3.cell(row=R_REVENUE, column=col).value not in (None, ""):
            last = col
    return col_letter(last)


# 'Financial Statements' rows, as written by dcf_comps_template's PL waterfall.
FS_R_YEARS = 4
FS_R_REVENUE = 6
FS_R_OP_INCOME = 11


def _hist_from_sheet(ws2):
    """Historical years / revenue / operating income off 'Financial Statements'."""
    years, revenue, op = [], [], []
    for col in range(3, 3 + 10):
        y = ws2.cell(row=FS_R_YEARS, column=col).value
        if y in (None, ""):
            break
        years.append(str(y))
        revenue.append(ws2.cell(row=FS_R_REVENUE, column=col).value)
        op.append(ws2.cell(row=FS_R_OP_INCOME, column=col).value)
    if op and str(ws2.cell(row=FS_R_OP_INCOME, column=2).value or "") != "Operating Income":
        raise SystemExit(
            f"ERROR: 'Financial Statements'!B{FS_R_OP_INCOME} is "
            f"{ws2.cell(row=FS_R_OP_INCOME, column=2).value!r}, not 'Operating Income' "
            f"- this workbook's layout does not match the template. Pass --op0 / "
            f"--peak-op / --peak-opm explicitly.")
    return years, revenue, op


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("xlsx")
    ap.add_argument("--op0", type=float, default=None,
                    help="Latest actual operating profit (JPY mn) - the ramp start")
    ap.add_argument("--peak-op", type=float, default=None,
                    help="Cycle-peak operating profit (JPY mn)")
    ap.add_argument("--peak-opm", type=float, default=None,
                    help="Cycle-peak operating margin (decimal)")
    ap.add_argument("--op0-label", default=None)
    ap.add_argument("--peak-label", default=None)
    ap.add_argument("--benchmark-ticker", default=None,
                    help="Comps ticker whose deal terms anchor Block E")
    ap.add_argument("--deal-note", action="append", default=None,
                    help="Free-text line(s) for the transaction-benchmark block")
    a = ap.parse_args()

    if not os.path.isfile(a.xlsx):
        sys.exit(f"ERROR: not found: {a.xlsx}")
    selftest_bn()

    wb = openpyxl.load_workbook(a.xlsx)
    for required in ("Executive Summary", "DCF Model", "Comps Analysis"):
        if required not in wb.sheetnames:
            sys.exit(f"ERROR: {a.xlsx} has no {required!r} sheet - not a DCF workbook "
                     f"from templates/dcf_comps_template.py")

    meta = _read_metadata(wb)
    ws1, ws2, ws3 = (wb["Executive Summary"], wb["Financial Statements"],
                     wb["DCF Model"])
    years, revenue, op = _hist_from_sheet(ws2)

    # 'Executive Summary'!C7 reads "5726.T (TSE Prime)"; taking it whole is what
    # put a doubled parenthesis in the old sheet title.
    ticker = str(ws1["C7"].value or "").split(" (")[0].strip()

    config = {
        "company_name": ws1["C6"].value or "",
        "ticker": ticker,
        "hist_years": years,
        "hist_revenue": revenue,
        "hist_operating_income": op,
        "base_year_revenue": ws3["C17"].value,
        "reverse_dcf": {k: v for k, v in {
            "op0": a.op0, "op0_label": a.op0_label,
            "peak_op": a.peak_op, "peak_label": a.peak_label,
            "peak_opm": a.peak_opm,
            "benchmark_ticker": a.benchmark_ticker,
            "deal_note": a.deal_note,
        }.items() if v is not None},
    }

    params, skip = resolve_params(config)
    if params is None:
        sys.exit(f"ERROR: cannot build the Reverse DCF sheet - {skip}\n"
                 f"       Pass --op0 / --peak-op / --peak-opm explicitly.")

    def _int(key, default):
        try:
            return int(str(meta.get(key, default)).strip())
        except (TypeError, ValueError):
            return default

    subject_row = _int("comps_subject_row", 0) or None
    bench_row = None
    if params.get("benchmark_ticker"):
        ws4 = wb["Comps Analysis"]
        want = str(params["benchmark_ticker"]).strip().upper().split(".")[0]
        for r in range(5, ws4.max_row + 1):
            got = str(ws4.cell(row=r, column=3).value or "").strip().upper().split(".")[0]
            if got and got == want:
                bench_row = r
                break
        if bench_row is None:
            print(f"  WARNING: benchmark ticker {params['benchmark_ticker']!r} is not "
                  f"in the comps table - Block E omitted.")

    build_reverse_dcf_sheet(wb, config, params, {
        "proj_last_col": _projection_last_col(ws3),
        "r_revenue": R_REVENUE,
        "r_da": R_DA,
        "r_ev_pgm": R_EV_PGM,
        "cmp_subject_row": subject_row,
        "cmp_stat_median_row": _int("comps_stat_median_row", 16),
        "cmp_benchmark_row": bench_row,
    })
    place_after(wb, "DCF Model")
    wb.save(a.xlsx)

    ws = wb[SHEET_NAME]
    n_f = sum(1 for row in ws.iter_rows() for c in row
              if isinstance(c.value, str) and c.value.startswith("="))
    print(f"Reverse DCF sheet written to {a.xlsx} ({n_f} live formulas)")
    print("NOTE: cached values were dropped on save -- run "
          "scripts/recalc_excel_com.py next, then scripts/validate_output.py.")


if __name__ == "__main__":
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    main()
