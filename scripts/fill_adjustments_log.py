"""fill_adjustments_log.py - write the analyst's Adjustments Log entries into a
generated workbook.

Per 手順書v2 §6-1 the log must carry ALL estimates, ALL design judgements and ALL
open items, including the ones still marked 未解決. Rows are inserted ABOVE the
'Pipeline Metadata' block (validate_output.py locates that block by its label,
not by row number), so the machine-readable metadata is never disturbed.

Nothing here is company-specific. The entries come from a config file:

    data/adjustments/<ticker>_adjustments.json   (auto-detected from the xlsx name)
    or --config PATH

Adding a ticker means adding a JSON file, never a new script — see
docs/DCFフォーマット標準メモ §3-5 and CLAUDE.md「銘柄コード入りスクリプトを作らない」.

Config schema:

    {
      "date": "2026-08-26",            // default 日付 for entries that omit one
      "entries": [
        ["DCF Model!C7", "Risk-Free Rate = 2.90%", "財務省 国債金利情報 …",
         "テンプレ既定 2.2%", "確定"],                       // 5 fields: date defaults
        ["2026-08-25", "DCF Model!C8", "Beta = 1.55", "…", "自動値 1.0", "確定"]
      ]                                                       // 6 fields: explicit date
    }

Columns are 日付 / セル / 変更内容 / 理由 / 元の値 / 状態. 状態 is one of
{確定 / 設計判断 / 推定 / 要確認 / 未解決 / DRAFT} (free text is allowed; these
are the values 手順書v2 defines).

Rule from tasks/lessons.md (2962): a cell value that starts with '=' is saved as
a FORMULA and can make Excel refuse to open the workbook. Every write below goes
through _put(), which asserts on that inside the loop - not on a sample.

openpyxl drops the cached formula values of every sheet on save, so run
scripts/recalc_excel_com.py on the workbook afterwards (and re-validate).
"""
import argparse
import json
import os
import re
import sys

import openpyxl
from openpyxl.styles import Font

BLACK_FONT = Font(name="Arial", size=10)
BOLD_FONT = Font(name="Arial", size=10, bold=True)

HEADERS = ("日付", "セル", "変更内容", "理由", "元の値", "状態")
FIRST_FREE_ROW = 6          # row 5 is the pipeline's own entry
ROW_HEIGHT = 30


def _put(ws, row, col, value, font=BLACK_FONT):
    s = "" if value is None else str(value)
    assert not s.startswith("="), (
        "cell R%dC%d would be stored as a formula: %r" % (row, col, s))
    c = ws.cell(row=row, column=col)
    c.value = value
    c.font = font
    return c


def normalise(entries, default_date):
    """Accept 5-field (date omitted) or 6-field rows; return 6-field tuples."""
    out = []
    for i, e in enumerate(entries):
        if isinstance(e, dict):
            e = [e.get(k) for k in ("date", "cell", "change", "reason",
                                    "previous", "status")]
            if e[0] is None:
                e[0] = default_date
        elif len(e) == len(HEADERS) - 1:
            e = [default_date] + list(e)
        elif len(e) != len(HEADERS):
            raise SystemExit(
                "ERROR: entry %d has %d field(s); expected %d (%s) or %d without 日付"
                % (i, len(e), len(HEADERS), " / ".join(HEADERS), len(HEADERS) - 1))
        if not e[0]:
            raise SystemExit("ERROR: entry %d has no 日付 and the config sets no "
                             "default \"date\"" % i)
        out.append(tuple(e))
    return out


def find_config(xlsx, explicit=None):
    if explicit:
        return explicit
    root = os.path.abspath(os.path.join(os.path.dirname(os.path.abspath(__file__)), ".."))
    m = re.match(r"([0-9A-Za-z]+)_", os.path.basename(xlsx))
    if not m:
        return None
    path = os.path.join(root, "data", "adjustments", "%s_adjustments.json" % m.group(1))
    return path if os.path.isfile(path) else None


def write_entries(wb, rows):
    ws = wb["Adjustments Log"]

    meta_row = None
    for r in range(1, ws.max_row + 1):
        v = ws.cell(row=r, column=2).value
        if isinstance(v, str) and v.startswith("Pipeline Metadata"):
            meta_row = r
            break
    if meta_row is None:
        raise SystemExit("Pipeline Metadata band not found - refusing to guess row "
                         "numbers (is this a workbook from dcf_comps_template?)")

    # Keep one blank row between the last entry and the metadata band.
    need = len(rows) - (meta_row - 1 - FIRST_FREE_ROW)
    if need > 0:
        ws.insert_rows(meta_row - 1, need)

    for i, rec in enumerate(rows):
        r = FIRST_FREE_ROW + i
        for j, val in enumerate(rec):
            _put(ws, r, 2 + j, val, font=BOLD_FONT if j == 1 else BLACK_FONT)
        ws.row_dimensions[r].height = ROW_HEIGHT
    return len(rows)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("xlsx")
    ap.add_argument("--config", default=None,
                    help="Entries JSON (default: data/adjustments/<ticker>_adjustments.json)")
    a = ap.parse_args()

    if not os.path.isfile(a.xlsx):
        sys.exit("ERROR: file not found: %s" % a.xlsx)
    cfg_path = find_config(a.xlsx, a.config)
    if not cfg_path or not os.path.isfile(cfg_path):
        sys.exit("ERROR: no adjustments config for %s.\n"
                 "       Create data/adjustments/<ticker>_adjustments.json (schema in "
                 "this file's docstring) or pass --config PATH.\n"
                 "       Do NOT copy this script to a ticker-specific name."
                 % os.path.basename(a.xlsx))
    with open(cfg_path, encoding="utf-8") as f:
        cfg = json.load(f)
    if not cfg.get("entries"):
        sys.exit("ERROR: %s has no 'entries'" % cfg_path)
    rows = normalise(cfg["entries"], cfg.get("date"))

    wb = openpyxl.load_workbook(a.xlsx)
    if "Adjustments Log" not in wb.sheetnames:
        sys.exit("ERROR: %s has no 'Adjustments Log' sheet" % a.xlsx)
    n = write_entries(wb, rows)
    wb.save(a.xlsx)

    print("Adjustments Log: wrote %d entries to %s (config: %s)"
          % (n, a.xlsx, os.path.relpath(cfg_path)))
    print("NOTE: cached values were dropped on save -- run "
          "scripts/recalc_excel_com.py next, then scripts/validate_output.py.")


if __name__ == "__main__":
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    main()
