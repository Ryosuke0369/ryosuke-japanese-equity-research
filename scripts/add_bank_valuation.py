"""
add_bank_valuation.py - Adds DDM and Residual Income sheets to a bank DCF workbook.

Banks break the standard UFCF/EV-net_debt DCF (debt is raw material, not funding;
NWC is undefined; regulatory capital constrains payouts). This script appends two
formula-driven sheets (values are never hardcoded, only inputs) to an already-generated
generate_dcf.py workbook:

  - "DDM": 2-stage Dividend Discount Model with a Ke x g sensitivity grid.
  - "Residual Income": Book-value rollforward + residual income model, plus a
    Justified P/B (ROE x Ke) grid.

It also patches the Comps Analysis sheet's percentile/median statistics formulas to
exclude the subject company's own row (row 5), which templates/dcf_comps_template.py
does not support natively (its stat range is always {col}5:{col}{last_row} inclusive
of whatever occupies row 5) -- required here because a bank's PER/PBR should not be
benchmarked against a median that includes itself (circular).

Usage:
    python scripts/add_bank_valuation.py models/8410_DCF_Model_<date>.xlsx
"""
import sys
import os
import openpyxl
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
from openpyxl.formatting.rule import FormulaRule
from openpyxl.utils import get_column_letter

# ── Style constants (mirrored from templates/dcf_comps_template.py for visual consistency) ──
BLACK_FONT  = Font(name="Arial", size=10, color="000000")
BLUE_FONT   = Font(name="Arial", size=10, color="0000FF")   # input cells
BOLD_FONT   = Font(name="Arial", size=10, bold=True)
TITLE_FONT  = Font(name="Arial", size=14, bold=True)
SUB_FONT    = Font(name="Arial", size=11, bold=True)
GREY_FONT   = Font(name="Arial", size=9, italic=True, color="808080")
HEADER_FONT = Font(name="Arial", size=11, bold=True, color="FFFFFF")

HEADER_FILL   = PatternFill(start_color="000080", end_color="000080", fill_type="solid")
LIGHT_GREEN   = PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid")
LIGHT_YELLOW  = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")
SUBTOTAL_FILL = PatternFill(start_color="F2F2F2", end_color="F2F2F2", fill_type="solid")
HIGHLIGHT_FILL = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")

THIN_BORDER = Border(left=Side(style="thin"), right=Side(style="thin"),
                      top=Side(style="thin"), bottom=Side(style="thin"))
_GRAY_SIDE = Side(style="thin", color="B0B0B0")
INPUT_BORDER = Border(left=_GRAY_SIDE, right=_GRAY_SIDE, top=_GRAY_SIDE, bottom=_GRAY_SIDE)
TOP_BOTTOM = Border(top=Side(style="thin"), bottom=Side(style="double"))
SECTION_BOTTOM = Border(bottom=Side(style="thin"))

FMT_YEN   = '#,##0;(#,##0)'
FMT_PCT   = '0.0%;(0.0%)'
FMT_PCT2  = '0.00%;(0.00%)'
FMT_RATIO = '0.00"x"'
FMT_INT   = '#,##0'
FMT_JPY_PS = '#,##0.0"円"'


def set_cell(ws, row, col, value, font=None, fmt=None, fill=None, border=None, alignment=None):
    c = ws.cell(row=row, column=col, value=value)
    if font: c.font = font
    if fmt: c.number_format = fmt
    if fill: c.fill = fill
    if border: c.border = border
    if alignment: c.alignment = alignment
    return c


def header_row(ws, row, col_start, labels):
    for i, lbl in enumerate(labels):
        c = ws.cell(row=row, column=col_start + i, value=lbl)
        c.font = HEADER_FONT
        c.fill = HEADER_FILL
        c.alignment = Alignment(horizontal="center", wrap_text=True)
        c.border = SECTION_BOTTOM


def section_title(ws, row, col, text, font=SUB_FONT):
    c = ws.cell(row=row, column=col, value=text)
    c.font = font
    return c


def col_letter(n):
    return get_column_letter(n)


# =====================================================================
# DDM SHEET
# =====================================================================
def build_ddm_sheet(wb):
    ws = wb.create_sheet("DDM")
    ws.sheet_properties.tabColor = "1F4E78"
    ws.column_dimensions["A"].width = 3
    ws.column_dimensions["B"].width = 28
    for col in "CDEFGHIJK":
        ws.column_dimensions[col].width = 13

    set_cell(ws, 2, 2, "DDM - 2-Stage Dividend Discount Model (PRIMARY METHOD - bank)", font=TITLE_FONT)
    set_cell(ws, 3, 2, "Standard UFCF DCF does not apply to a bank (see 'DCF Model' sheet, used only as an auxiliary ATM-fee-franchise cross-check).",
             font=GREY_FONT)

    # ── Inputs ──
    section_title(ws, 5, 2, "Inputs")
    rows = [
        ("Risk-Free Rate",       "='DCF Model'!C7",  FMT_PCT2),
        ("Beta",                 "='DCF Model'!C8",  '0.00'),
        ("Equity Risk Premium",  "='DCF Model'!C9",  FMT_PCT2),
        ("Cost of Equity (Ke)",  "='DCF Model'!C23", FMT_PCT2),
        ("Shares Outstanding",   "='DCF Model'!C15", FMT_INT),
        ("Current Price (JPY)",  "='Executive Summary'!C9", FMT_YEN),
        ("Terminal Growth (g)",  0.015,               FMT_PCT2),
    ]
    r = 6
    for label, val, fmt in rows:
        set_cell(ws, r, 2, label, font=BOLD_FONT)
        is_input = isinstance(val, (int, float))
        set_cell(ws, r, 3, val, font=(BLUE_FONT if is_input else BLACK_FONT), fmt=fmt,
                 border=(INPUT_BORDER if is_input else None))
        r += 1
    KE_CELL = "$C$9"
    G_CELL = "$C$12"
    SHARES_CELL = "$C$10"
    PRICE_CELL = "$C$11"

    set_cell(ws, r + 1, 2, "Guard (Ke must exceed g)", font=BOLD_FONT)
    set_cell(ws, r + 1, 3, f'=IF({KE_CELL}<={G_CELL},"ERROR: Ke <= g, model invalid","OK")',
             font=BLACK_FONT)
    guard_row = r + 1

    # ── DPS projection (Base scenario) ──
    dps_title_row = guard_row + 2
    section_title(ws, dps_title_row, 2, "Dividend per Share Forecast (Base scenario)")
    yr_row = dps_title_row + 1
    header_row(ws, yr_row, 3, ["FY2027(E)", "FY2028(E)", "FY2029(E)", "FY2030(E)", "FY2031(E)"])

    dps_row = yr_row + 1
    set_cell(ws, dps_row, 2, "DPS (JPY)", font=BOLD_FONT)
    dps_values = [11.00, 11.50, 12.00, 12.50, 13.00]
    for i, v in enumerate(dps_values):
        set_cell(ws, dps_row, 3 + i, v, font=BLUE_FONT, fmt='0.00', border=INPUT_BORDER)
    set_cell(ws, dps_row, 8, "FY2027(E)=11.00 matches company guidance (dividend forecast, flat vs FY2026/3 actual)", font=GREY_FONT)

    df_row = dps_row + 1
    set_cell(ws, df_row, 2, "Discount Factor", font=BOLD_FONT)
    for i in range(5):
        col = 3 + i
        set_cell(ws, df_row, col, f"=1/(1+{KE_CELL})^{i+1}", font=BLACK_FONT, fmt='0.0000')

    pv_row = df_row + 1
    set_cell(ws, pv_row, 2, "PV of DPS", font=BOLD_FONT)
    for i in range(5):
        col = 3 + i
        cl = col_letter(col)
        set_cell(ws, pv_row, col, f"={cl}{dps_row}*{cl}{df_row}", font=BLACK_FONT, fmt='0.00')

    # ── Valuation ──
    val_title_row = pv_row + 2
    section_title(ws, val_title_row, 2, "Valuation")
    r = val_title_row + 1
    set_cell(ws, r, 2, "Sum of PV of DPS (Yr 1-5)", font=BOLD_FONT)
    set_cell(ws, r, 3, f"=SUM({col_letter(3)}{pv_row}:{col_letter(7)}{pv_row})", font=BLACK_FONT, fmt='0.00')
    sum_pv_row = r
    r += 1
    set_cell(ws, r, 2, "Terminal Value (Yr 5)", font=BOLD_FONT)
    set_cell(ws, r, 3, f"={col_letter(7)}{dps_row}*(1+{G_CELL})/({KE_CELL}-{G_CELL})", font=BLACK_FONT, fmt='0.00')
    tv_row = r
    r += 1
    set_cell(ws, r, 2, "PV of Terminal Value", font=BOLD_FONT)
    set_cell(ws, r, 3, f"=C{tv_row}*{col_letter(7)}{df_row}", font=BLACK_FONT, fmt='0.00')
    pv_tv_row = r
    r += 1
    set_cell(ws, r, 2, "Implied Value per Share (JPY)", font=BOLD_FONT, fill=LIGHT_GREEN, border=TOP_BOTTOM)
    set_cell(ws, r, 3, f"=ROUND(C{sum_pv_row}+C{pv_tv_row},0)", font=BOLD_FONT, fmt=FMT_YEN, fill=LIGHT_GREEN, border=TOP_BOTTOM)
    implied_row = r
    r += 1
    set_cell(ws, r, 2, "Upside / (Downside) vs Current Price", font=BOLD_FONT)
    set_cell(ws, r, 3, f"=C{implied_row}/{PRICE_CELL}-1", font=BLACK_FONT, fmt=FMT_PCT)

    # ── Sensitivity: Ke (rows) x g (cols) ──
    sens_title_row = r + 2
    section_title(ws, sens_title_row, 2, "Sensitivity: Ke (rows) x Terminal g (cols)")
    ke_vals = [0.052 + 0.0025 * i for i in range(9)]   # 5.2% .. 7.2%, 0.25% step
    g_vals = [0.005 + 0.0025 * i for i in range(9)]    # 0.5% .. 2.5%, 0.25% step

    hdr_row = sens_title_row + 1
    set_cell(ws, hdr_row, 2, "Ke \\ g", font=BOLD_FONT)
    for j, g in enumerate(g_vals):
        set_cell(ws, hdr_row, 3 + j, g, font=BOLD_FONT, fmt=FMT_PCT2, alignment=Alignment(horizontal="center"))

    dps_range = f"${col_letter(3)}${dps_row}:${col_letter(7)}${dps_row}"
    for i, ke in enumerate(ke_vals):
        row = hdr_row + 1 + i
        set_cell(ws, row, 2, ke, font=BOLD_FONT, fmt=FMT_PCT2)
        ke_cell = f"$B{row}"
        for j, g in enumerate(g_vals):
            col = 3 + j
            g_cell = f"{col_letter(col)}${hdr_row}"
            formula = (
                f'=IF({ke_cell}<={g_cell},"",'
                f'ROUND(SUMPRODUCT({dps_range},1/(1+{ke_cell})^{{1,2,3,4,5}})'
                f'+{col_letter(7)}{dps_row}*(1+{g_cell})/({ke_cell}-{g_cell})/(1+{ke_cell})^5,0))'
            )
            set_cell(ws, row, col, formula, font=BLACK_FONT, fmt=FMT_YEN, border=THIN_BORDER)

    ws.freeze_panes = "C6"
    return {"ke_cell": KE_CELL, "g_cell": G_CELL, "shares_cell": SHARES_CELL,
            "price_cell": PRICE_CELL, "dps_row": dps_row, "implied_row": implied_row,
            "dps_col_start": 3, "dps_col_end": 7}


# =====================================================================
# RESIDUAL INCOME SHEET
# =====================================================================
def build_ri_sheet(wb, ddm_refs, bank_book_value_mn, net_income_actual_mn):
    ws = wb.create_sheet("Residual Income")
    ws.sheet_properties.tabColor = "1F4E78"
    ws.column_dimensions["A"].width = 3
    ws.column_dimensions["B"].width = 30
    for col in "CDEFGHIJK":
        ws.column_dimensions[col].width = 13

    set_cell(ws, 2, 2, "Residual Income Model (PRIMARY METHOD - bank)", font=TITLE_FONT)
    set_cell(ws, 3, 2, "Value = Book Value + PV(Residual Income). When ROE approaches Ke, value converges to book value.",
              font=GREY_FONT)

    section_title(ws, 5, 2, "Inputs")
    r = 6
    set_cell(ws, r, 2, "Beginning Book Value (JPY mn)", font=BOLD_FONT)
    set_cell(ws, r, 3, bank_book_value_mn, font=BLUE_FONT, fmt=FMT_YEN, border=INPUT_BORDER)
    bv0_row = r
    r += 1
    set_cell(ws, r, 2, "Shares Outstanding", font=BOLD_FONT)
    set_cell(ws, r, 3, f"=DDM!{ddm_refs['shares_cell'][1:]}", font=BLACK_FONT, fmt=FMT_INT)
    shares_row = r
    r += 1
    set_cell(ws, r, 2, "BPS_0 (JPY)", font=BOLD_FONT)
    set_cell(ws, r, 3, f"=C{bv0_row}/C{shares_row}*1000000", font=BLACK_FONT, fmt='0.0"円"')
    bps0_row = r
    r += 1
    set_cell(ws, r, 2, "Cost of Equity (Ke)", font=BOLD_FONT)
    set_cell(ws, r, 3, f"=DDM!{ddm_refs['ke_cell'][1:]}", font=BLACK_FONT, fmt=FMT_PCT2)
    ke_row = r
    r += 1
    set_cell(ws, r, 2, "Terminal Growth (g)", font=BOLD_FONT)
    set_cell(ws, r, 3, f"=DDM!{ddm_refs['g_cell'][1:]}", font=BLACK_FONT, fmt=FMT_PCT2)
    g_row = r
    KE_REF, G_REF = f"$C${ke_row}", f"$C${g_row}"

    # ── ROE forecast ──
    r += 2
    section_title(ws, r, 2, "ROE Forecast (Base scenario)")
    yr_row = r + 1
    header_row(ws, yr_row, 3, ["FY2027(E)", "FY2028(E)", "FY2029(E)", "FY2030(E)", "FY2031(E)"])
    roe_row = yr_row + 1
    set_cell(ws, roe_row, 2, "ROE", font=BOLD_FONT)
    roe_values = [0.061, 0.065, 0.068, 0.070, 0.070]
    for i, v in enumerate(roe_values):
        set_cell(ws, roe_row, 3 + i, v, font=BLUE_FONT, fmt=FMT_PCT, border=INPUT_BORDER)
    set_cell(ws, roe_row, 8, "FY2027(E)=6.1% = company net income guidance 17,000 / beginning equity 280,491", font=GREY_FONT)

    # ── Rollforward ──
    r = roe_row + 2
    section_title(ws, r, 2, "Book Value Rollforward (JPY mn)")
    bv_hdr = r + 1
    set_cell(ws, bv_hdr, 2, "", font=BOLD_FONT)
    for i in range(5):
        set_cell(ws, bv_hdr, 3 + i, ["FY2027(E)", "FY2028(E)", "FY2029(E)", "FY2030(E)", "FY2031(E)"][i],
                 font=BOLD_FONT, alignment=Alignment(horizontal="center"))

    ni_row = bv_hdr + 1
    div_row = ni_row + 1
    bv_row = div_row + 1
    ri_row = bv_row + 1
    pv_ri_row = ri_row + 1

    set_cell(ws, ni_row, 2, "Net Income_t = ROE_t x BV_(t-1)", font=BOLD_FONT)
    set_cell(ws, div_row, 2, "Dividends_t (from DDM)", font=BOLD_FONT)
    set_cell(ws, bv_row, 2, "Ending Book Value", font=BOLD_FONT)
    set_cell(ws, ri_row, 2, "Residual Income_t = (ROE_t - Ke) x BV_(t-1)", font=BOLD_FONT)
    set_cell(ws, pv_ri_row, 2, "PV of Residual Income", font=BOLD_FONT)

    ddm_dps_row = ddm_refs["dps_row"]
    for i in range(5):
        col = 3 + i
        cl = col_letter(col)
        prev_bv_cell = f"C{bv_row}" if i == 0 else f"{col_letter(col-1)}{bv_row}"
        prev_bv_ref = f"$C${bv0_row}" if i == 0 else prev_bv_cell
        set_cell(ws, ni_row, col, f"={cl}{roe_row}*{prev_bv_ref}", font=BLACK_FONT, fmt=FMT_YEN)
        set_cell(ws, div_row, col, f"=DDM!{cl}{ddm_dps_row}*C{shares_row}/1000000", font=BLACK_FONT, fmt=FMT_YEN)
        set_cell(ws, bv_row, col, f"={prev_bv_ref}+{cl}{ni_row}-{cl}{div_row}", font=BLACK_FONT, fmt=FMT_YEN, border=THIN_BORDER)
        set_cell(ws, ri_row, col, f"=({cl}{roe_row}-{KE_REF})*{prev_bv_ref}", font=BLACK_FONT, fmt=FMT_YEN)
        set_cell(ws, pv_ri_row, col, f"={cl}{ri_row}/(1+{KE_REF})^{i+1}", font=BLACK_FONT, fmt=FMT_YEN)

    # ── Valuation ──
    r = pv_ri_row + 2
    section_title(ws, r, 2, "Valuation")
    r += 1
    set_cell(ws, r, 2, "Sum of PV of Residual Income", font=BOLD_FONT)
    set_cell(ws, r, 3, f"=SUM(C{pv_ri_row}:G{pv_ri_row})", font=BLACK_FONT, fmt=FMT_YEN)
    sum_pv_ri_row = r
    r += 1
    set_cell(ws, r, 2, "Terminal Residual Income", font=BOLD_FONT)
    set_cell(ws, r, 3, f"=G{ri_row}*(1+{G_REF})/({KE_REF}-{G_REF})", font=BLACK_FONT, fmt=FMT_YEN)
    term_ri_row = r
    r += 1
    set_cell(ws, r, 2, "PV of Terminal Residual Income", font=BOLD_FONT)
    set_cell(ws, r, 3, f"=C{term_ri_row}/(1+{KE_REF})^5", font=BLACK_FONT, fmt=FMT_YEN)
    pv_term_ri_row = r
    r += 1
    set_cell(ws, r, 2, "Equity Value (JPY mn)", font=BOLD_FONT, fill=LIGHT_GREEN, border=TOP_BOTTOM)
    set_cell(ws, r, 3, f"=C{bv0_row}+C{sum_pv_ri_row}+C{pv_term_ri_row}", font=BOLD_FONT, fmt=FMT_YEN, fill=LIGHT_GREEN, border=TOP_BOTTOM)
    ev_row = r
    r += 1
    set_cell(ws, r, 2, "Implied Value per Share (JPY)", font=BOLD_FONT, fill=LIGHT_GREEN, border=TOP_BOTTOM)
    set_cell(ws, r, 3, f"=ROUND(C{ev_row}/C{shares_row}*1000000,0)", font=BOLD_FONT, fmt=FMT_YEN, fill=LIGHT_GREEN, border=TOP_BOTTOM)
    implied_row = r
    r += 1
    set_cell(ws, r, 2, "Upside / (Downside) vs Current Price", font=BOLD_FONT)
    set_cell(ws, r, 3, f"=C{implied_row}/DDM!{ddm_refs['price_cell'][1:]}-1", font=BLACK_FONT, fmt=FMT_PCT)
    r += 2
    set_cell(ws, r, 2,
             "Note: ROE sits close to Ke, so Residual Income per year is small and the implied value sits close to "
             "BPS_0 -- the theoretical price is structurally anchored to book value, not to growth assumptions. "
             "Whether ROE durably exceeds Ke is the entire valuation question for this stock.",
             font=GREY_FONT)

    # ── Justified P/B grid ──
    r += 2
    section_title(ws, r, 2, "Justified P/B = (ROE - g) / (Ke - g)")
    r += 1
    roe_grid_vals = [0.045 + 0.005 * i for i in range(9)]   # 4.5% .. 8.5%, 0.5% step
    ke_grid_vals = [0.052 + 0.005 * i for i in range(5)]    # 5.2% .. 7.2%, 0.5% step

    hdr_row2 = r
    set_cell(ws, hdr_row2, 2, "ROE \\ Ke", font=BOLD_FONT)
    for j, ke in enumerate(ke_grid_vals):
        set_cell(ws, hdr_row2, 3 + j, ke, font=BOLD_FONT, fmt=FMT_PCT2, alignment=Alignment(horizontal="center"))

    pb_first_data_row = hdr_row2 + 1
    for i, roe in enumerate(roe_grid_vals):
        row = pb_first_data_row + i
        set_cell(ws, row, 2, roe, font=BOLD_FONT, fmt=FMT_PCT)
        roe_cell = f"$B{row}"
        for j, ke in enumerate(ke_grid_vals):
            col = 3 + j
            ke_cell = f"{col_letter(col)}${hdr_row2}"
            formula = f'=IF({ke_cell}<={G_REF},"",ROUND(({roe_cell}-{G_REF})/({ke_cell}-{G_REF}),2))'
            set_cell(ws, row, col, formula, font=BLACK_FONT, fmt=FMT_RATIO, border=THIN_BORDER)
    pb_last_data_row = pb_first_data_row + len(roe_grid_vals) - 1

    # Conditional formatting: highlight cells within 0.05x of the current P/B (Market_Cap/Book_Value)
    cf_range = f"C{pb_first_data_row}:{col_letter(2+len(ke_grid_vals))}{pb_last_data_row}"
    current_pb_ref = f"DDM!{ddm_refs['price_cell'][1:]}*C{shares_row}/1000000/C{bv0_row}"
    rule = FormulaRule(
        formula=[f'AND(ISNUMBER({col_letter(3)}{pb_first_data_row}),ABS({col_letter(3)}{pb_first_data_row}-({current_pb_ref}))<0.05)'],
        fill=HIGHLIGHT_FILL,
    )
    # openpyxl FormulaRule anchors to the top-left cell of the range; apply per-cell via a relative formula instead
    rule2 = FormulaRule(formula=[f'ABS(C{pb_first_data_row}-({current_pb_ref}))<0.05'], fill=HIGHLIGHT_FILL)
    ws.conditional_formatting.add(cf_range, rule2)

    r = pb_last_data_row + 2
    set_cell(ws, r, 2, "Implied Share Price at each (ROE, Ke) node = BPS_0 x Justified P/B (JPY)", font=BOLD_FONT)
    r += 1
    hdr_row3 = r
    set_cell(ws, hdr_row3, 2, "ROE \\ Ke", font=BOLD_FONT)
    for j, ke in enumerate(ke_grid_vals):
        set_cell(ws, hdr_row3, 3 + j, ke, font=BOLD_FONT, fmt=FMT_PCT2, alignment=Alignment(horizontal="center"))
    price_first_row = hdr_row3 + 1
    for i, roe in enumerate(roe_grid_vals):
        row = price_first_row + i
        pb_row_ref = pb_first_data_row + i
        set_cell(ws, row, 2, roe, font=BOLD_FONT, fmt=FMT_PCT)
        for j, ke in enumerate(ke_grid_vals):
            col = 3 + j
            cl = col_letter(col)
            formula = f'=IF({cl}{pb_row_ref}="","",ROUND({cl}{pb_row_ref}*$C${bps0_row},0))'
            set_cell(ws, row, col, formula, font=BLACK_FONT, fmt=FMT_YEN, border=THIN_BORDER)

    ws.freeze_panes = "C6"
    return {"bv0_row": bv0_row, "bps0_row": bps0_row, "implied_row": implied_row}


# =====================================================================
# COMPS ANALYSIS PATCH: exclude subject company (row 5) from stats
# =====================================================================
def patch_comps_exclude_self(wb):
    if "Comps Analysis" not in wb.sheetnames:
        print("  [add_bank_valuation] No 'Comps Analysis' sheet found -- skipping self-exclusion patch.")
        return False
    ws = wb["Comps Analysis"]

    # Find the last comp data row by scanning column B (Company) from row 5 downward
    last_comp_row = 4
    row = 5
    while ws.cell(row=row, column=2).value not in (None, ""):
        last_comp_row = row
        row += 1

    if last_comp_row <= 5:
        print("  [add_bank_valuation] Only one (or zero) comp rows found -- nothing to exclude, skipping patch.")
        return False

    stat_rows = {"25th Percentile": None, "Median (50th)": None, "75th Percentile": None}
    for r in range(1, 40):
        label = ws.cell(row=r, column=2).value
        if label in stat_rows:
            stat_rows[label] = r

    stat_col_map = [(4, 10), (5, 11), (6, 12), (7, 13), (8, 14), (9, 15)]  # (dst_col, src_col)
    patched = 0
    for label, r in stat_rows.items():
        if r is None:
            continue
        for dst_col, src_col in stat_col_map:
            src_letter = col_letter(src_col)
            rng = f"{src_letter}6:{src_letter}{last_comp_row}"  # start at row 6, skip subject row 5
            if label == "25th Percentile":
                formula = f"=PERCENTILE({rng},0.25)"
            elif label == "Median (50th)":
                formula = f"=MEDIAN({rng})"
            else:
                formula = f"=PERCENTILE({rng},0.75)"
            ws.cell(row=r, column=dst_col, value=formula)
            patched += 1
    print(f"  [add_bank_valuation] Patched {patched} Comps Analysis stat formulas to exclude subject row 5 "
          f"(stats now over rows 6:{last_comp_row}).")
    return True


def main():
    if len(sys.argv) < 2:
        print("Usage: python scripts/add_bank_valuation.py <path-to-dcf-model.xlsx>")
        sys.exit(1)
    path = sys.argv[1]
    if not os.path.exists(path):
        print(f"ERROR: file not found: {path}")
        sys.exit(1)

    wb = openpyxl.load_workbook(path)

    if "DDM" in wb.sheetnames:
        del wb["DDM"]
    if "Residual Income" in wb.sheetnames:
        del wb["Residual Income"]

    ddm_refs = build_ddm_sheet(wb)
    build_ri_sheet(wb, ddm_refs, bank_book_value_mn=280491, net_income_actual_mn=13476)
    patch_comps_exclude_self(wb)

    wb.save(path)
    print(f"Saved bank valuation sheets (DDM, Residual Income) into: {path}")
    print("Run scripts/recalc_excel_com.py on this file next to compute formula values.")


if __name__ == "__main__":
    main()
