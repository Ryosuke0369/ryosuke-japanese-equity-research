"""add_segment_bridge.py - append a 'Segment Bridge' sheet to a DCF workbook.

Segment data is disclosed for too few comparable periods to feed the main DCF's
`segments` override block (which also switches COGS to a segment-EBIT
back-solve), so it lives on its own sheet: segment revenue and profit, bridged
to the consolidated figures with the corporate/elimination adjustment shown and
any tie-out variance surfaced rather than forced to zero.

Nothing here is company-specific. The numbers come from a config file:

    data/segments/<ticker>_segments.json      (auto-detected from the xlsx name)
    or --config PATH

which is the only thing that changes per ticker. Adding a ticker means adding a
JSON file, never a new script — see docs/DCFフォーマット標準メモ §3-5 and
CLAUDE.md「銘柄コード入りスクリプトを作らない」.

Usage:
    python scripts/add_segment_bridge.py models/3110_DCF_Model_20260822.xlsx
    python scripts/add_segment_bridge.py <xlsx> --config data/segments/3687_segments.json

Config schema (all sections optional except `periods` and `segments`):

    {
      "title":   "Segment Bridge - Nitto Boseki (3110)",
      "intro":   ["one line per paragraph above the first table"],
      "periods": ["FY2025/3", "FY2026/3", "Q1 FY2027/3"],
      "segments": [
        {"label": "Electronic Materials  電子材料",
         "revenue": [52093, 61418, 18082],
         "profit":  [13880, 19391, 7036],
         "da":      [null, 7229, null],
         "note":    "optional per-row note (column F)"}
      ],
      "adjustment": {
        "revenue": "derive",            // "derive" = consolidated - segment total
        "revenue_note": "...",
        "profit": [-858, -2160, -486],
        "profit_label": "Adjustment  調整額 (全社費用・消去)",
        "da": [null, 259, null]
      },
      "consolidated": {"revenue": [...], "profit": [...], "da": [...],
                       "revenue_note": "...", "profit_note": "...", "da_note": "..."},
      "show_margin_table": true,
      "show_ebitda_row": true,          // consolidated OP + consolidated D&A
      "extra_tables": [
        {"title": "Customer Concentration",
         "headers": ["Customer", "FY2024/9", "% of Rev"],
         "formats": ["text", "yen", "pct"],
         "rows": [["Kioxia", 1613, {"formula": "=C{row}/7995"}],
                  ["Total", {"formula": "=SUM(C{first}:C{last})"},
                            {"formula": "=C{row}/7995"}]],
         "data_rows": 1,                // rows counted by {first}:{last} (default: all)
         "note": "..."}
      ],
      "notes": ["free text under the tables"],
      "open_items": [["2026-08-22", "'DCF Model'!C5", "what changed",
                      "why", "previous value", "未解決"]]
    }

openpyxl drops the cached formula values of every sheet on save, so run
scripts/recalc_excel_com.py on the workbook afterwards (and re-validate).
"""
import argparse
import json
import os
import re
import sys

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

SHEET_NAME = "Segment Bridge"

TITLE_FONT = Font(name="Arial", size=14, bold=True, color="1F4E79")
SUB_FONT = Font(name="Arial", size=11, bold=True, color="1F4E79")
BOLD_FONT = Font(name="Arial", size=10, bold=True)
BLACK_FONT = Font(name="Arial", size=10)
BLUE_FONT = Font(name="Arial", size=10, color="0070C0")
GREY_FONT = Font(name="Arial", size=9, italic=True, color="808080")
HEADER_FONT = Font(name="Arial", size=11, bold=True, color="FFFFFF")

HEADER_FILL = PatternFill(start_color="000080", end_color="000080", fill_type="solid")
LIGHT_GREEN = PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid")
LIGHT_YELLOW = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")

SECTION_BOTTOM = Border(bottom=Side(style="thin"))
TOP_BOTTOM = Border(top=Side(style="thin"), bottom=Side(style="double"))
_GREY = Side(style="thin", color="B0B0B0")
INPUT_BORDER = Border(left=_GREY, right=_GREY, top=_GREY, bottom=_GREY)

FMT = {
    "yen": "#,##0;(#,##0)",
    "pct": "0.1%;(0.1%)",
    "ratio": "0.00x",
    "int": "#,##0",
    "text": None,
}

FIRST_DATA_COL = 3      # column C
NOTE_COL_OFFSET = 1     # note column sits one past the last period column


# --------------------------------------------------------------------- helpers
def put(ws, row, col, value, font=None, fmt=None, fill=None, border=None, align=None):
    c = ws.cell(row=row, column=col)
    c.value = value
    if font:
        c.font = font
    if fmt:
        c.number_format = fmt
    if fill:
        c.fill = fill
    if border:
        c.border = border
    if align:
        c.alignment = align
    return c


def label(ws, row, col, text, font=BLACK_FONT, fill=None, border=None):
    """Write a TEXT label.

    A string that starts with '=' is stored by openpyxl as a FORMULA and can make
    Excel refuse to open the workbook entirely (tasks/lessons.md, 2962). Assert
    on every write, not on a hand-picked sample.
    """
    assert not str(text).startswith("="), (
        "label at R%dC%d starts with '=': %r" % (row, col, text))
    return put(ws, row, col, text, font=font, fill=fill, border=border)


def section_title(ws, row, text, last_col):
    label(ws, row, 2, text, font=SUB_FONT)
    for c in range(2, last_col + 1):
        ws.cell(row=row, column=c).border = SECTION_BOTTOM


def header_row(ws, row, values, fill=HEADER_FILL):
    for i, v in enumerate(values):
        label(ws, row, 2 + i, v, font=HEADER_FONT, fill=fill)


def cl(col):
    return openpyxl.utils.get_column_letter(col)


# ---------------------------------------------------------------- sheet blocks
class _Builder:
    def __init__(self, ws, cfg):
        self.ws = ws
        self.cfg = cfg
        self.periods = cfg["periods"]
        self.n = len(self.periods)
        self.last_col = FIRST_DATA_COL + self.n - 1
        self.note_col = self.last_col + NOTE_COL_OFFSET
        self.r = 1

    # -- small helpers ----------------------------------------------------
    def _values_row(self, row, values, font, fmt, fill=None, border=None):
        for i in range(self.n):
            v = values[i] if values and i < len(values) else None
            if v is None:
                continue
            put(self.ws, row, FIRST_DATA_COL + i, v, font=font, fmt=fmt,
                fill=fill, border=border)

    def _formula_row(self, row, template, font, fmt, fill=None, border=None,
                     cols=None):
        """template(col_letter, i) -> formula string. `cols` limits the periods.

        A period no segment discloses (3110 gives segment D&A for FY2026/3 only)
        gets no cell at all: SUM over blanks would print a confident 0 where the
        honest answer is "not disclosed".
        """
        for i in (range(self.n) if cols is None else cols):
            put(self.ws, row, FIRST_DATA_COL + i,
                template(cl(FIRST_DATA_COL + i), i),
                font=font, fmt=fmt, fill=fill, border=border)

    def _note(self, row, text):
        if text:
            label(self.ws, row, self.note_col, text, font=GREY_FONT)

    def _has(self, key):
        return any(s.get(key) for s in self.cfg["segments"])

    # -- one metric block (revenue / profit / D&A) ------------------------
    def block(self, n, key, title, total_label, adj_label, cons_label,
              derive_adjustment=False, show_variance=False):
        """Segment rows -> total -> adjustment -> consolidated (+ variance)."""
        segs = [s for s in self.cfg["segments"] if s.get(key)]
        if not segs:
            return None
        adj = self.cfg.get("adjustment") or {}
        cons = (self.cfg.get("consolidated") or {}).get(key)
        active = [i for i in range(self.n)
                  if any(i < len(s[key]) and s[key][i] is not None for s in segs)]

        section_title(self.ws, self.r, "%d. %s" % (n, title), self.note_col)
        self.r += 1
        header_row(self.ws, self.r, ["Segment"] + list(self.periods) + ["Note"])
        self.r += 1

        first = self.r
        for s in segs:
            label(self.ws, self.r, 2, s["label"])
            self._values_row(self.r, s[key], BLUE_FONT, FMT["yen"])
            self._note(self.r, s.get("note") if key == "revenue" else None)
            self.r += 1
        last = self.r - 1

        label(self.ws, self.r, 2, total_label, font=BOLD_FONT)
        self._formula_row(self.r, lambda c, i: "=SUM(%s%d:%s%d)" % (c, first, c, last),
                          BOLD_FONT, FMT["yen"], border=SECTION_BOTTOM, cols=active)
        total_row = self.r
        self.r += 1

        adj_row = None
        adj_vals = adj.get(key)
        if derive_adjustment and adj_vals in (None, "derive") and cons:
            label(self.ws, self.r, 2, adj_label)
            self._formula_row(
                self.r,
                lambda c, i: ("=%s-%s%d" % (cons[i], c, total_row)
                              if i < len(cons) and cons[i] is not None else None),
                BLACK_FONT, FMT["yen"])
            self._note(self.r, adj.get(key + "_note"))
            adj_row = self.r
            self.r += 1
        elif isinstance(adj_vals, list):
            label(self.ws, self.r, 2, adj_label)
            self._values_row(self.r, adj_vals, BLUE_FONT, FMT["yen"])
            self._note(self.r, adj.get(key + "_note"))
            adj_row = self.r
            self.r += 1

        bridged_row = None
        if show_variance and adj_row:
            label(self.ws, self.r, 2, "Bridged total  計算値", font=BOLD_FONT)
            self._formula_row(self.r,
                              lambda c, i: "=%s%d+%s%d" % (c, total_row, c, adj_row),
                              BOLD_FONT, FMT["yen"], cols=active)
            bridged_row = self.r
            self.r += 1

        cons_row = None
        if cons:
            label(self.ws, self.r, 2, cons_label, font=BOLD_FONT)
            self._values_row(self.r, cons, BOLD_FONT, FMT["yen"],
                             fill=LIGHT_GREEN,
                             border=None if show_variance else TOP_BOTTOM)
            self._note(self.r, (self.cfg.get("consolidated") or {}).get(key + "_note"))
            cons_row = self.r
            self.r += 1

        if bridged_row and cons_row:
            label(self.ws, self.r, 2, "Variance  差異 (計算値 - 開示値)", font=BOLD_FONT)
            self._formula_row(
                self.r,
                lambda c, i: ("=%s%d-%s%d" % (c, bridged_row, c, cons_row)
                              if i < len(cons) and cons[i] is not None else None),
                BOLD_FONT, FMT["yen"], fill=LIGHT_YELLOW, border=TOP_BOTTOM)
            self._note(self.r, "Rounding of the disclosed JPY mn figures. A non-zero "
                               "variance is shown rather than forced to zero.")
            self.r += 1

        self.r += 1
        return {"first": first, "last": last, "total": total_row,
                "adjustment": adj_row, "consolidated": cons_row,
                "rows": {s["label"]: first + i for i, s in enumerate(segs)},
                "order": [s["label"] for s in segs]}


def build_sheet(wb, cfg):
    if SHEET_NAME in wb.sheetnames:
        del wb[SHEET_NAME]
    ws = wb.create_sheet(SHEET_NAME)
    ws.sheet_properties.tabColor = "7030A0"

    b = _Builder(ws, cfg)
    ws.column_dimensions["A"].width = 3
    ws.column_dimensions["B"].width = 36
    for i in range(b.n):
        ws.column_dimensions[cl(FIRST_DATA_COL + i)].width = 16
    ws.column_dimensions[cl(b.note_col)].width = 46

    b.r = 2
    label(ws, b.r, 2, cfg.get("title", SHEET_NAME), font=TITLE_FONT)
    b.r += 1
    for line in cfg.get("intro", []):
        label(ws, b.r, 2, line, font=GREY_FONT)
        b.r += 1
    b.r += 1

    n = 1
    rev = b.block(n, "revenue", "Segment Revenue (JPY mn)",
                  "Segment total  セグメント計",
                  (cfg.get("adjustment") or {}).get(
                      "revenue_label", "Adjustment  調整額 (内部売上高消去等)"),
                  "Consolidated revenue  連結売上高",
                  derive_adjustment=True)
    if rev:
        n += 1

    prof = b.block(n, "profit", "Segment Profit to Consolidated Operating Profit (JPY mn)",
                   "Segment profit total  セグメント利益計",
                   (cfg.get("adjustment") or {}).get(
                       "profit_label", "Adjustment  調整額 (全社費用・消去)"),
                   "Consolidated operating profit  連結営業利益 (開示値)",
                   show_variance=True)
    if prof:
        n += 1

    # ── Segment operating margin (derived, live off the two blocks above) ──
    if cfg.get("show_margin_table", True) and rev and prof:
        section_title(ws, b.r, "%d. Segment Operating Margin" % n, b.note_col)
        b.r += 1
        header_row(ws, b.r, ["Segment"] + list(b.periods) + ["Note"])
        b.r += 1
        # Matched by label: a segment with profit but no disclosed revenue
        # (3687's "Other (incl. CVC)") simply has no margin row, rather than
        # being silently paired with whichever row happens to sit at the same
        # offset in the other block.
        for name in rev["order"]:
            if name not in prof["rows"]:
                continue
            label(ws, b.r, 2, name)
            for i in range(b.n):
                c = cl(FIRST_DATA_COL + i)
                put(ws, b.r, FIRST_DATA_COL + i,
                    '=IFERROR(%s%d/%s%d,"N/A")'
                    % (c, prof["rows"][name], c, rev["rows"][name]),
                    font=BLACK_FONT, fmt=FMT["pct"])
            b.r += 1
        label(ws, b.r, 2, "Consolidated  連結", font=BOLD_FONT)
        for i in range(b.n):
            c = cl(FIRST_DATA_COL + i)
            put(ws, b.r, FIRST_DATA_COL + i,
                '=IFERROR(%s%d/%s%d,"N/A")'
                % (c, prof["consolidated"] or prof["total"],
                   c, rev["consolidated"] or rev["total"]),
                font=BOLD_FONT, fmt=FMT["pct"], fill=LIGHT_GREEN, border=TOP_BOTTOM)
        b.r += 2
        n += 1

    da = b.block(n, "da", "Depreciation & Amortisation by Segment (JPY mn)",
                 "Segment D&A total  計",
                 (cfg.get("adjustment") or {}).get("da_label", "Adjustment  調整額"),
                 "Consolidated D&A  連結減価償却費 (開示値)",
                 show_variance=True)
    if da:
        n += 1
        if cfg.get("show_ebitda_row", True) and prof and prof["consolidated"]:
            label(ws, b.r - 1, 2, "Consolidated EBITDA  (連結営業利益 + 連結D&A)",
                  font=BOLD_FONT)
            da_row = da["consolidated"] or da["total"]
            for i in range(b.n):
                c = cl(FIRST_DATA_COL + i)
                if ws.cell(row=da_row, column=FIRST_DATA_COL + i).value is None:
                    continue
                put(ws, b.r - 1, FIRST_DATA_COL + i,
                    '=IFERROR(%s%d+%s%d,"N/A")'
                    % (c, prof["consolidated"], c, da_row),
                    font=BOLD_FONT, fmt=FMT["yen"], fill=LIGHT_GREEN)
            label(ws, b.r - 1, b.note_col,
                  "This is the core_ebitda basis used on 'Comps Analysis'.",
                  font=GREY_FONT)
            b.r += 1

    # ── Free-form tables (customer concentration, normalisation bridges, ...) ──
    # Formula cells are written as {"formula": "..."} and may reference:
    #   {row} {first} {last}      this table's current / first / last data row
    #   {r0} {r1} ...             this table's own rows, in order
    #   {c0} {c1} ...             the period columns (C, D, ...)
    #   {revenue_r0} {profit_r0} {da_r0} ...   the n-th segment row of a block
    #   {revenue_total} {profit_total} {da_total}
    #   {revenue_cons}  {profit_cons}  {da_cons}
    # so a normalisation bridge stays live off the blocks above instead of
    # restating their numbers as constants.
    anchors = {}
    for key, blk in (("revenue", rev), ("profit", prof), ("da", da)):
        if not blk:
            continue
        for i, name in enumerate(blk["order"]):
            anchors["%s_r%d" % (key, i)] = blk["rows"][name]
        anchors["%s_total" % key] = blk["total"]
        anchors["%s_cons" % key] = blk["consolidated"] or blk["total"]
        anchors["%s_adj" % key] = blk["adjustment"] or blk["total"]
    for i in range(b.n):
        anchors["c%d" % i] = cl(FIRST_DATA_COL + i)

    for tbl in cfg.get("extra_tables", []):
        section_title(ws, b.r, "%d. %s" % (n, tbl["title"]), b.note_col)
        n += 1
        b.r += 1
        header_row(ws, b.r, tbl["headers"])
        b.r += 1
        rows = tbl["rows"]
        n_data = tbl.get("data_rows", len(rows))
        first, last = b.r, b.r + n_data - 1
        own = {"r%d" % i: first + i for i in range(len(rows))}
        fmts = tbl.get("formats") or []
        for row_vals in rows:
            for j, v in enumerate(row_vals):
                col = 2 + j
                default_fmt = fmts[j] if j < len(fmts) else "text"
                if isinstance(v, dict):
                    fmt = FMT.get(v.get("fmt", default_fmt))
                    put(ws, b.r, col,
                        v["formula"].format(row=b.r, first=first, last=last,
                                            **own, **anchors),
                        font=BOLD_FONT if v.get("bold") else BLACK_FONT, fmt=fmt,
                        fill=LIGHT_GREEN if v.get("highlight") else None)
                elif isinstance(v, (int, float)):
                    put(ws, b.r, col, v, font=BLUE_FONT, fmt=FMT.get(default_fmt),
                        border=INPUT_BORDER)
                elif v is None:
                    put(ws, b.r, col, "-", font=GREY_FONT,
                        align=Alignment(horizontal="right"))
                else:
                    label(ws, b.r, col, v, font=BOLD_FONT if j == 0 else BLACK_FONT)
            b.r += 1
        if tbl.get("note"):
            label(ws, b.r, 2, tbl["note"], font=GREY_FONT)
            ws.row_dimensions[b.r].height = 45
            b.r += 1
        b.r += 1

    if cfg.get("notes"):
        label(ws, b.r, 2, "Notes", font=SUB_FONT)
        b.r += 1
        for note in cfg["notes"]:
            label(ws, b.r, 2, note, font=GREY_FONT)
            b.r += 1

    ws.freeze_panes = "C%d" % (FIRST_DATA_COL + 3)
    return ws


# ------------------------------------------------------------- adjustments log
def append_open_items(wb, items):
    """Append open items to the 'Adjustments Log' manual-entry table."""
    if not items:
        return 0
    if "Adjustments Log" not in wb.sheetnames:
        print("  WARNING: no 'Adjustments Log' sheet - open items not recorded.")
        return 0
    ws = wb["Adjustments Log"]
    row = 5
    while ws.cell(row=row, column=2).value not in (None, ""):
        if str(ws.cell(row=row, column=2).value).startswith("Pipeline Metadata"):
            break
        row += 1
    for entry in items:
        for j, v in enumerate(entry):
            label(ws, row, 2 + j, v)
        row += 1
    return len(items)


# ------------------------------------------------------------------------ main
def find_config(xlsx, explicit=None):
    if explicit:
        return explicit
    root = os.path.abspath(os.path.join(os.path.dirname(os.path.abspath(__file__)), ".."))
    m = re.match(r"([0-9A-Za-z]+)_", os.path.basename(xlsx))
    if not m:
        return None
    path = os.path.join(root, "data", "segments", "%s_segments.json" % m.group(1))
    return path if os.path.isfile(path) else None


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("xlsx")
    ap.add_argument("--config", default=None,
                    help="Segment config JSON (default: data/segments/<ticker>_segments.json)")
    a = ap.parse_args()

    if not os.path.isfile(a.xlsx):
        sys.exit("ERROR: file not found: %s" % a.xlsx)
    cfg_path = find_config(a.xlsx, a.config)
    if not cfg_path or not os.path.isfile(cfg_path):
        sys.exit("ERROR: no segment config for %s.\n"
                 "       Create data/segments/<ticker>_segments.json (schema in this "
                 "file's docstring) or pass --config PATH.\n"
                 "       Do NOT copy this script to a ticker-specific name."
                 % os.path.basename(a.xlsx))
    with open(cfg_path, encoding="utf-8") as f:
        cfg = json.load(f)
    for required in ("periods", "segments"):
        if not cfg.get(required):
            sys.exit("ERROR: %s: '%s' is required" % (cfg_path, required))

    wb = openpyxl.load_workbook(a.xlsx)
    build_sheet(wb, cfg)
    n_items = append_open_items(wb, cfg.get("open_items"))
    wb.save(a.xlsx)

    print("Segment Bridge written to %s (config: %s)"
          % (a.xlsx, os.path.relpath(cfg_path)))
    if n_items:
        print("  Appended %d open item(s) to the Adjustments Log." % n_items)
    print("NOTE: cached values were dropped on save -- run "
          "scripts/recalc_excel_com.py next, then scripts/validate_output.py.")


if __name__ == "__main__":
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    main()
