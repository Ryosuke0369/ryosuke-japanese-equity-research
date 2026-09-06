"""Append the reverse-comps sheet ('Implied Multiple Analysis') to an existing
market-analysis workbook, positioned right after 'Implied Growth Analysis'.

Reuses templates/market_analysis_template._build_reverse_comps_sheet so the
sheet is identical to what generate_market_analysis_excel would produce. Pulls
the Comps distribution from the same DCF Excel via extract_dcf_data.

    python scripts/append_reverse_comps_sheet.py 4192

NOTE: openpyxl drops cached formula values of the other sheets on save, so run
      scripts/recalc_excel_com.py on the workbook afterwards to restore them.
"""
import sys
import os
import argparse

import openpyxl

sys.path.insert(0, os.path.join(os.path.dirname(os.path.abspath(__file__)), '..', 'templates'))
from market_analysis_template import extract_dcf_data, _build_reverse_comps_sheet  # noqa: E402

PROJECT_ROOT = os.path.abspath(os.path.join(os.path.dirname(os.path.abspath(__file__)), '..'))

# Per-ticker config for the reverse-comps sheet (current price + targets).
CONFIGS = {
    '4192': {
        'ticker': '4192.T',
        'company_name': 'SpiderPlus & Co.',
        'current_price': 293,            # integrated report snapshot (2026-06-03)
        'forward_revenue': 5900,
        'price_targets': {'損切り': 225, '現在': 293, '第1利確': 564, '第2利確': 750},
    },
}


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('ticker', nargs='?', default='4192')
    ap.add_argument('--date', default='20260603')
    ap.add_argument('--xlsx', default=None)
    ap.add_argument('--dcf', default=None)
    args = ap.parse_args()

    cfg = CONFIGS.get(args.ticker)
    if cfg is None:
        raise SystemExit(f"No reverse-comps config for ticker {args.ticker}; add one to CONFIGS.")

    xlsx_path = args.xlsx or os.path.join(
        PROJECT_ROOT, 'reports', f'{args.ticker}_market_analysis_{args.date}.xlsx')
    dcf_path = args.dcf or os.path.join(
        PROJECT_ROOT, 'models', f'{args.ticker}_DCF_Model_{args.date}.xlsx')
    for p in (xlsx_path, dcf_path):
        if not os.path.exists(p):
            raise FileNotFoundError(p)

    dd = extract_dcf_data(dcf_path)
    if not (dd.get('comps') or {}).get('present'):
        raise SystemExit(f"{dcf_path} has no usable Comps Analysis data; nothing to append.")

    wb = openpyxl.load_workbook(xlsx_path)
    if 'Implied Multiple Analysis' in wb.sheetnames:
        del wb['Implied Multiple Analysis']
    _build_reverse_comps_sheet(wb, cfg, dd)

    # Position right after 'Implied Growth Analysis' (逆算同士を隣接).
    name = 'Implied Multiple Analysis'
    if 'Implied Growth Analysis' in wb.sheetnames:
        target = wb.sheetnames.index('Implied Growth Analysis') + 1
        wb.move_sheet(name, offset=target - wb.sheetnames.index(name))

    wb.save(xlsx_path)
    print(f"Sheets now: {wb.sheetnames}")
    print(f"Saved: {os.path.normpath(xlsx_path)}")
    print("Run: python scripts/recalc_excel_com.py "
          f"reports/{args.ticker}_market_analysis_{args.date}.xlsx  (restore formula caches)")


if __name__ == '__main__':
    try:
        sys.stdout.reconfigure(encoding='utf-8')
    except Exception:
        pass
    main()
