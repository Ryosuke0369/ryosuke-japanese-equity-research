"""arbitration.py - 追補6 §X: which of the two DCF legs survives when they disagree.

Why this lives in scripts/ and not batch/
-----------------------------------------
This is pipeline logic, not batch orchestration. It used to sit in
`batch/arbitrate_divergence.py` and run as a separate manual step after
generation, which meant every regeneration silently threw the arbitration away:
`generate_dcf.py` knows nothing about it, so a fresh workbook goes straight back
to averaging two legs that 追補5 forbids averaging. 6857 アドバンテスト's Target
read 4,288 (PGM alone) before the フェーズ2 regeneration and 8,021 (the midpoint)
after it — not a valuation change, just the arbitration being absent.

Moving the rule engine here lets `generate_dcf.py` finish the job itself
(追補12 §A-3). `batch/arbitrate_divergence.py` and `batch/demote_dcf_leg.py` are
now thin command-line wrappers over these functions, so the batch-side scan and
the in-pipeline application can never drift apart on what the rule says.

The ladder (unchanged from 追補6 §X)
------------------------------------
    divergence > 3.0x (multiples OR prices)  ->  the midpoint average is forbidden

    1. growth-premium regime (§Q: assumed > 25x AND pgm-implied < 10x)
       -> demote Exit, Target = PGM alone (conservative floor).
    2. trough signature (§U: latest-FY OPM <= window p25), EXCEPT on a strictly
       monotonic series - a trend is not a trough, and 追補6 §Y forbids treating
       one with mean reversion.
       -> re-generate with mid-cycle normalisation FIRST, then fall through to 3.
    3. normal regime -> OBSERVED-MULTIPLE BAND TEST against the peers' own
       EV/EBITDA 25th-75th percentiles. Exactly one leg outside -> demote it.

Every function here takes explicit paths rather than resolving a ticker code
against a basis date, because the pipeline knows its own output file and the
batch scan does not.
"""
import csv
import json
import os

BAND_RATIO_MAX = 1.5   # 追補7 §AA — above this the peer band is too wide to anchor on
LEVERAGE_MAX = 0.6     # 追補8 §AD-2 — above this net_debt/EV the point estimate is meaningless
LOW_WACC = 0.05        # 追補8 §AG — below this WACC the beta clamp + leverage mix inflates Target
DIVERGENCE_MAX = 3.0   # 追補6 §X — above this the midpoint average is forbidden

RULE_TAG = {"exit": "追補4 §Q", "pgm": "追補5 §U", "x": "追補6 §X"}
LEG_PREFIX = {"exit": "DCF - Exit Multiple", "pgm": "DCF - Perpetuity Growth"}
TREATED = {"exit": "追補4 §Q", "pgm": "追補5 §U", "pgm6": "追補6 §X"}

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))


# =====================================================================
# company type — decides whether the DCF arbitration applies at all
# =====================================================================
# 追補12 §A-3: the arbitration decides between two DCF legs, so it is meaningful
# only for a model whose Target IS a DCF. 型D (banks) values on DDM + Residual
# Income and a DCF does not hold for them at all; running a DCF-leg arbitration
# there would arbitrate between two numbers nobody uses.
#
# The type is NOT inferred. 8410's overrides carry net_debt=0 and
# base_year_ar/inv/ap=0, which looks like a bank fingerprint but is really just
# the mechanism for forcing ΔNWC to zero - a non-bank can legitimately have the
# same shape. Guessing the type from that would be exactly the silent inference
# this pipeline forbids, so it comes from an explicit `company_type` override.
DCF_TYPES = ("A", "B", "C", "E")   # E = the non-bank part is a DCF (SOTP later)
NON_DCF_TYPES = ("D",)             # D = DDM / Residual Income is the主手法


def resolve_company_type(overrides):
    """'A'..'E' from the overrides, or None when not declared."""
    if not overrides:
        return None
    v = overrides.get("company_type")
    if isinstance(v, str) and v.strip().upper() in DCF_TYPES + NON_DCF_TYPES:
        return v.strip().upper()
    return None


def arbitration_applies(company_type):
    """(bool, reason). Never silent: the reason is printed by the caller."""
    if company_type in NON_DCF_TYPES:
        return False, (f"型{company_type}（銀行）は DCF が成立せず Target は DDM+RI が主手法。"
                       f"DCF 2脚の裁定は適用対象外")
    if company_type is None:
        return True, ("company_type 未宣言 — DCF 型（A/B/C/E）とみなして適用する。"
                      "型D なら overrides に company_type: \"D\" を明示すること")
    return True, f"型{company_type}（DCF 型）— 適用対象"


# =====================================================================
# inputs, read from the workbook and its side files
# =====================================================================
def pct(sorted_vals, q):
    """Excel PERCENTILE.INC — linear interpolation between order statistics."""
    if len(sorted_vals) == 1:
        return sorted_vals[0]
    i = (len(sorted_vals) - 1) * q
    lo = int(i)
    hi = min(lo + 1, len(sorted_vals) - 1)
    return sorted_vals[lo] + (sorted_vals[hi] - sorted_vals[lo]) * (i - lo)


def divergence(xlsx):
    """(pgm_implied, assumed, divergence) computed FROM THE WORKBOOK, or None.

    This used to be scraped out of `<xlsx>_validation.txt` with a regex, which
    made the arbitration impossible to run before validate_output.py — the exact
    ordering 追補12 §A-3 asks for. The three numbers are cached cells, so they
    are read directly, with the same arithmetic validate's check 11 uses:

        implied = terminal value (PGM) / Year-5 EBITDA
        gap     = max(implied, assumed) / min(implied, assumed)

    The workbook must have been recalculated; without cached values this returns
    None and the caller reports why rather than guessing.
    """
    try:
        import openpyxl
        from templates.dcf_comps_template import R_TV_PGM, R_YR5_EBITDA
    except ImportError:
        return None
    if not os.path.exists(xlsx):
        return None
    try:
        ws = openpyxl.load_workbook(xlsx, data_only=True)["DCF Model"]
    except Exception:
        return None

    def num(v):
        return v if isinstance(v, (int, float)) and not isinstance(v, bool) else None

    tv = num(ws.cell(row=R_TV_PGM, column=3).value)
    ebitda5 = num(ws.cell(row=R_YR5_EBITDA, column=3).value)
    assumed = num(ws["C14"].value)
    if not tv or not ebitda5 or not assumed:
        return None
    implied = tv / ebitda5
    return implied, assumed, max(implied, assumed) / min(implied, assumed)


def peer_band(comps_csv):
    """(p25, p75, n, [values]) of PEER EV/EBITDA — the subject row is excluded.

    Same arithmetic as batch/exit_multiple.py so the band and the exit multiple
    can never disagree about what a peer multiple is.
    """
    if not comps_csv or not os.path.exists(comps_csv):
        return None
    rows = list(csv.DictReader(open(comps_csv, encoding="utf-8")))
    if len(rows) < 2:
        return None

    def num(r, k):
        try:
            return float(r[k])
        except (TypeError, ValueError, KeyError):
            return None

    vals = []
    for r in rows[1:]:                       # row 0 is the subject
        e, m, n = num(r, "EBITDA"), num(r, "Market_Cap"), num(r, "Net_Debt")
        if e and m is not None and n is not None and e > 0:
            vals.append((m + n) / e)
    if len(vals) < 3:
        return None
    vals.sort()
    return (pct(vals, 0.25), pct(vals, 0.75), len(vals), vals)


def opm_series(overrides_path=None, cache_path=None):
    """(latest OPM, window p25, full series) from overrides first, then the cache."""
    if overrides_path and os.path.exists(overrides_path):
        d = json.load(open(overrides_path, encoding="utf-8"))
        rev, oi = d.get("hist_revenue"), d.get("hist_operating_income")
        if rev and oi and len(rev) == len(oi):
            s = [o / r for r, o in zip(rev, oi) if r]
            if len(s) >= 3:
                return s[-1], pct(sorted(s), 0.25), s
    if cache_path and os.path.exists(cache_path):
        d = json.load(open(cache_path, encoding="utf-8"))
        inc = d.get("income", {})
        ys = sorted({k for v in inc.values() for k in v})
        s = []
        for y in ys:
            r = inc.get("Total Revenue", {}).get(y)
            o = inc.get("Operating Income", {}).get(y)
            if r and o is not None:
                s.append(o / r)
        if len(s) >= 3:
            return s[-1], pct(sorted(s), 0.25), s
    return None


def strictly_monotonic(series):
    """'increasing' / 'decreasing' / None - the same test batch/pregen_check.py uses."""
    if not series or len(series) < 3:
        return None
    if all(b > a for a, b in zip(series, series[1:])):
        return "increasing"
    if all(b < a for a, b in zip(series, series[1:])):
        return "decreasing"
    return None


def midcycle_recorded(overrides_path):
    """True when the overrides record that a mid-cycle normalisation was applied.

    追補5 §U says "re-generate with mid-cycle normalisation FIRST, then re-judge".
    `already_treated()` cannot answer that: it looks for a demoted leg, which is
    the OUTCOME of the re-judgement, not the input to it. The analyst records the
    normalisation in the overrides' `_scenario_note`, so that is what is read.
    """
    if not overrides_path or not os.path.exists(overrides_path):
        return False
    try:
        d = json.load(open(overrides_path, encoding="utf-8"))
    except (OSError, ValueError):
        return False
    txt = " ".join(str(v) for k, v in d.items() if k.startswith("_"))
    return ("ミッドサイクル正常化" in txt) or ("§U" in txt)


def legs_and_wacc(xlsx, overrides_path=None):
    """(pgm_price, exit_price, price, shares, wacc, net_debt) from the workbook."""
    if not os.path.exists(xlsx):
        return None
    try:
        import openpyxl
        wb = openpyxl.load_workbook(xlsx, data_only=True)
        e, d = wb["Executive Summary"], wb["DCF Model"]
    except Exception:
        return None
    pgm, ex, price = e.cell(16, 3).value, e.cell(17, 3).value, e.cell(9, 3).value
    wacc = None
    for r in range(1, 40):
        lab = d.cell(r, 2).value
        if isinstance(lab, str) and lab.startswith("WACC") and isinstance(d.cell(r, 3).value, float):
            wacc = d.cell(r, 3).value
            break
    nd = sh = None
    if overrides_path and os.path.exists(overrides_path):
        ov = json.load(open(overrides_path, encoding="utf-8"))
        nd, sh = ov.get("net_debt"), ov.get("shares_outstanding")
    return pgm, ex, price, sh, wacc, nd


def already_treated(xlsx):
    """Which rule (if any) has already been stamped into the Executive Summary."""
    if not os.path.exists(xlsx):
        return None
    try:
        import openpyxl
        ws = openpyxl.load_workbook(xlsx)["Executive Summary"]
    except Exception:
        return None
    for r in range(1, 40):
        v = ws.cell(r, 2).value
        if not isinstance(v, str):
            continue
        for leg, tag in TREATED.items():
            if tag in v and "参考" in v:
                return ("exit" if "Exit" in v else "pgm", tag)
    return None


# =====================================================================
# the ladder
# =====================================================================
def arbitrate(xlsx, overrides_path=None, comps_csv=None, cache_path=None,
              code=None, divergence_fallback=None):
    """Apply 追補6 §X to one workbook. Returns the verdict dict, or None.

    `divergence_fallback` lets a caller supply (pgm_implied, assumed, div) when
    the workbook has no cached values — batch/arbitrate_divergence.py passes the
    numbers parsed out of an old validation report so pre-フェーズ2 models can
    still be scanned.
    """
    c = divergence(xlsx) or divergence_fallback
    if not c:
        return None
    pgm_i, assumed, div = c
    out = dict(code=code, pgm_implied=pgm_i, assumed=assumed, div=div,
               regime="-", verdict="-", demote=None, band=None, opm=None,
               band_ratio=None, band_unstable=False, price_div=None, pgm_price=None,
               exit_price=None, lev=None, lev_mkt=None, wacc=None, low_wacc=False, na=False,
               treated=already_treated(xlsx))

    lw = legs_and_wacc(xlsx, overrides_path)
    if lw:
        pgm_p, ex_p, price, sh, wacc, nd = lw
        out["wacc"] = wacc
        # 追補8 §AG — a very low WACC is its own review flag, independent of divergence
        if wacc is not None and wacc < LOW_WACC:
            out["low_wacc"] = True
        if all(isinstance(x, (int, float)) for x in (pgm_p, ex_p)) and min(pgm_p, ex_p) > 0:
            out["price_div"] = max(pgm_p, ex_p) / min(pgm_p, ex_p)
            out["pgm_price"], out["exit_price"] = pgm_p, ex_p
            if nd is not None and sh:
                # 追補8 §AD-2 の分母は「本モデルの株式価値」であって時価総額ではない。
                # 論点は『EV のわずかな差が株式価値で数倍に増幅される』ことなので、
                # 増幅の分母になるのはモデルが出した株式価値のほう。
                eq_model = (pgm_p + ex_p) / 2 * sh / 1e6
                ev_model = eq_model + nd
                out["lev"] = (nd / ev_model) if ev_model else None
                out["lev_mkt"] = (nd / (price * sh / 1e6 + nd)) if price else None

    pdiv = out.get("price_div") or 1.0
    # 追補8 §AD — the trigger is an OR: multiples OR prices
    if div <= DIVERGENCE_MAX and pdiv <= DIVERGENCE_MAX:
        out["regime"] = "収斂"
        out["verdict"] = "裁定不要（倍率乖離・株価乖離とも ≤ 3.0倍）"
        return out

    if div <= DIVERGENCE_MAX and pdiv > DIVERGENCE_MAX:
        # 追補8 §AD-2 — leverage-amplified. Neither leg is wrong; the equity value is
        # simply a small difference between two large numbers, so a point estimate is
        # not meaningful. The band test cannot separate the legs (the multiples agree).
        out["regime"] = "レバレッジ増幅型(§AD-2)"
        lev = out.get("lev")
        if lev is not None and lev > LEVERAGE_MAX:
            out["demote"] = None
            out["na"] = True
            out["verdict"] = ("**Target = N/A（高レバレッジ・判定不能）** + 要目視レビュー"
                              "（net_debt/EV = %.2f > %.1f）" % (lev, LEVERAGE_MAX))
        else:
            out["verdict"] = ("保守側の脚を採用 + EV差異原因をサマリーに分析記録"
                              "（net_debt/EV = %s ≤ %.1f）"
                              % ("%.2f" % lev if lev is not None else "算出不可", LEVERAGE_MAX))
        return out

    # 1. growth-premium regime
    if assumed > 25.0 and pgm_i < 10.0:
        out["regime"] = "成長プレミアム(§Q)"
        out["demote"] = "exit"
        out["verdict"] = "Exit降格 → Target = PGM 単独"
        return out

    # 2. trough signature
    o = opm_series(overrides_path, cache_path)
    out["opm"] = o
    if o and o[0] <= o[1]:
        # §U asks "is the denominator sitting in a cyclical trough". Its signature
        # (latest <= p25) is ALSO true of every strictly monotonic decline, where
        # the latest value is the minimum by construction - and 追補6 §Y says in so
        # many words that a trend must not be treated with mean reversion. Left
        # unqualified, §U therefore captured 6273 / 6367 / 6506 (OPM 31.31 -> 25.26
        # -> 24.02 -> 22.62, 9.47 -> 8.92 -> 8.45 -> 8.27, 12.29 -> 11.50 -> 9.33
        # -> 8.73) and parked them at "re-generate with mid-cycle normalisation",
        # which is the one treatment §Y forbids for them. A trend is not a trough:
        # those names go to the band test like any other normal-regime name.
        mono = strictly_monotonic(o[2] if len(o) > 2 else None)
        if mono:
            out["regime"] = "単調(%s)・§U非該当" % ("増加" if mono == "increasing" else "減少")
        else:
            out["regime"] = "トラフ(§U)"
            # falls through to the band test only once mid-cycle regeneration is done
            if not (out["treated"] or midcycle_recorded(overrides_path)):
                out["verdict"] = "ミッドサイクル正常化で再生成 → 再判定"
                return out
            if not out["treated"]:
                out["regime"] = "トラフ(§U)・正常化済"

    # 3. normal regime — observed-multiple band test
    b = peer_band(comps_csv)
    if not b:
        out["regime"] = out["regime"] if out["regime"] != "-" else "通常"
        out["verdict"] = "Peer倍率バンドを作れず → 要手動レビュー"
        return out
    p25, p75, n, _ = b
    out["band"] = (p25, p75, n)
    out["band_ratio"] = (p75 / p25) if p25 else None
    pgm_out = pgm_i < p25 or pgm_i > p75
    exit_out = assumed < p25 or assumed > p75
    if out["regime"] == "-":
        out["regime"] = "通常"
    if pgm_out and not exit_out:
        out["demote"] = "pgm"
        out["verdict"] = "PGM降格 → Target = Exit 単独"
    elif exit_out and not pgm_out:
        out["demote"] = "exit"
        out["verdict"] = "Exit降格 → Target = PGM 単独"
    elif not pgm_out and not exit_out:
        out["verdict"] = "両脚バンド内 → 平均維持 + 要目視レビュー"
    else:
        out["verdict"] = "両脚バンド外 → 要手動レビューキュー"

    # 追補7 §AA — the band itself can be a weak anchor. Flag ONLY the names that
    # were resolved BY the band test (a wide band is only a problem when we leant
    # on it), so the §Q growth-premium names are deliberately out of scope.
    if out["demote"] and out["band_ratio"] and out["band_ratio"] > BAND_RATIO_MAX:
        out["band_unstable"] = True
        out["verdict"] += "（**バンド不安定・要目視レビュー**: p75/p25 = %.2f > %.1f）" % (
            out["band_ratio"], BAND_RATIO_MAX)
    return out


# =====================================================================
# applying the verdict to the workbook
# =====================================================================
def _find_row(ws, prefix, col=2, limit=40):
    for r in range(1, min(ws.max_row, limit) + 1):
        v = ws.cell(r, col).value
        if isinstance(v, str) and v.startswith(prefix):
            return r
    return None


def demotion_note(leg, rule, div, pgm_implied, assumed, band=""):
    """The Adjustments-Log sentence for a demotion. Text only, no side effects."""
    if rule == "x":
        kept = "Exit" if leg == "pgm" else "PGM"
        return (
            "【追補6 §X】PGMとExitの乖離が **%s倍** で 3.0倍を超えたため、"
            "**乖離そのものを裁定トリガーとする追補6 §X により中点平均を禁止**し、どちらの脚を落とすかを判定した。"
            "■**観測倍率バンドテスト**: Peer の実測 EV/EBITDA の第1〜第3四分位バンドは **%s倍**。"
            "これに対し **PGM逆算 %s倍 / 仮定Exit %s倍**。"
            "**%s法のみがバンドの外にあるため %s法を[参考]に降格し、Target = %s 単独**とした。"
            "■バンドテストは市場倍率を無条件のアンカーにするものではない（それではスクリーナーが市場の複写に堕する）。"
            "**実際に取引されているレンジの外に出た脚だけを取り除く**という位置づけである。"
            % (div, band or "—", pgm_implied, assumed,
               "PGM" if leg == "pgm" else "Exit",
               "PGM" if leg == "pgm" else "Exit", kept))
    if leg == "exit":
        return (
            "【追補4 §Q】Exit法を[参考]に降格し Target = PGM 単独とした。"
            "PGM逆算Exit倍率 %s倍 に対し仮定Exit倍率(Peer中央値) %s倍 で乖離 %s倍。"
            "永久成長率1.0%%固定の下ではPGM逆算Exitは5〜7倍にしかならず、"
            "Peer中央値が織り込む長期成長と構造的に両立しないため、両者の平均は"
            "『相容れない2つの世界観の中点』にすぎない。"
            "PGM単独Targetは保守フロアであり、現値との差は市場が織り込む成長プレミアムとして読むこと。"
            % (pgm_implied, assumed, div))
    return (
        "【追補5 §U】型Bミッドサイクル正常化を適用して再生成した後もPGMとExitの乖離が %s倍 残ったため、"
        "§U の手順に従いどちらを正とするかを個別判断し、**PGM法を[参考]に降格して Target = Exit 単独**とした。"
        "PGM逆算のターミナルEV/EBITDAは %s倍 で、上場する事業会社が実際に取引される倍率の下限を大きく下回り、"
        "ターミナル価値として経験的に成立しない。一方の仮定Exit倍率 %s倍(Peer中央値)は市場で観測される水準の内側にある。"
        "片方の脚が観測可能な範囲の外に出た場合は、内側にある脚を正とする。"
        "乖離の原因は本銘柄固有の洞察ではなく、永久成長率1.0%%固定の単段階モデルが"
        "資本集約型(capex/D&A が持続的に1.5倍超)を評価しきれないという既知のモデル上の限界である"
        "(バッチ後の2段階成長モデル導入課題として登録済み)。"
        % (div, pgm_implied, assumed))


def apply_demotion(xlsx, leg, div, pgm_implied, assumed, band="", rule="auto",
                   reason="", quiet=False):
    """Demote ONE leg to a reference row so Target = the other leg alone.

    Following 手順書§5-5 (never write a label starting with '='):
      1. the demoted leg's label -> marked [参考・Target不算入 — <rule>]
      2. C10 target -> =IF(ISNUMBER(Cn),ROUND(Cn,0),"N/A") pointing at the
         SURVIVING leg. Still a formula and still free of the comps rows, so
         validate check 14 stays out of FAIL and reports WARN "does not average
         C16:C17" — the intended signal that a leg was deliberately dropped.
      3. B20 note -> the ticker-specific reason is appended.

    The workbook is saved by openpyxl, which does not compute: the caller MUST
    recalculate before validating, or every value-level check reports SKIP (and
    フェーズ2 #8 turns SKIP into FAIL).
    """
    import openpyxl
    wb = openpyxl.load_workbook(xlsx)
    ws = wb["Executive Summary"]
    r_dem = _find_row(ws, LEG_PREFIX[leg])
    r_keep = _find_row(ws, LEG_PREFIX["pgm" if leg == "exit" else "exit"])
    r_note = _find_row(ws, "Note: Target Mid")
    r_tgt = _find_row(ws, "Target Price")
    if not (r_dem and r_keep and r_note and r_tgt):
        raise ValueError("Exec Summary rows not found — refusing to guess row numbers.")

    tag_rule = RULE_TAG["x"] if rule == "x" else RULE_TAG[leg]
    tag = "[参考・Target不算入 — %s]" % tag_rule
    label = ws.cell(r_dem, 2).value
    if tag not in label:
        ws.cell(r_dem, 2).value = label + " " + tag
    assert not str(ws.cell(r_dem, 2).value).startswith("="), "label must not start with '='"

    # each methodology row carries its implied value in column C of the SAME row
    # (B16/C16 = PGM, B17/C17 = Exit), so the surviving leg's value cell is C<r_keep>.
    ws.cell(r_tgt, 3).value = '=IF(ISNUMBER(C%d),ROUND(C%d,0),"N/A")' % (r_keep, r_keep)

    sentence = demotion_note(leg, rule, div, pgm_implied, assumed, band)
    if reason:
        sentence += reason
    cur = ws.cell(r_note, 2).value
    if tag_rule not in cur:
        ws.cell(r_note, 2).value = cur + " ■" + sentence
    wb.save(xlsx)
    if not quiet:
        print("demoted %s leg in %s" % (leg.upper(), xlsx))
        print("  B%d = %s" % (r_dem, ws.cell(r_dem, 2).value))
        print("  C%d = %s" % (r_tgt, ws.cell(r_tgt, 3).value))
        print("NOTE: cached values dropped — recalculate before validating.")
    return {"leg": leg, "rule": tag_rule, "target_row": r_tgt,
            "surviving_row": r_keep, "demoted_row": r_dem}
