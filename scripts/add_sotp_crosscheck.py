"""
add_sotp_crosscheck.py - Adds a SOTP (sum-of-the-parts) cross-check sheet to the
AEON (8267) consolidated DCF workbook.

AEON is a holding company for many listed and unlisted subsidiaries; a consolidated
DCF is an approximation and SOTP is the theoretically correct lens (see prompt).
This script appends one sheet, "SOTP Cross-Check", built entirely from Excel formulas
(inputs are hardcoded market data / ownership stakes as of the research date; all
downstream math is formula-driven so changing an input recalculates the sheet).

IMPORTANT FINDING baked into this script (see notes written into the sheet itself):
Of the 9 subsidiaries originally in scope for SOTP, THREE went private in 2025
(AEON Mall 2025-06-27, AEON Delight 2025-07-17, Welcia Holdings merged into Tsuruha
2025-11-27) -- after the source prompt was authored. Only 6 remain listed with a
market price: AEON Financial Service, Tsuruha, AEON Kyushu, AEON Hokkaido, Maxvalu
Tokai, Ministop. Per the prompt's own fallback instruction ("非上場事業は算入せず、
上場子会社持分のみのフロア値として提示"), this SOTP presents the listed-stake sum as
an explicit FLOOR value, not a full sum-of-the-parts -- it does not attempt to price
the now-wholly-owned AEON Mall/AEON Delight or the GMS/SM retail core, since no
reliable non-guessed EBITDA breakdown for those units was obtained.

Usage:
    python scripts/add_sotp_crosscheck.py models/8267_DCF_Model_<date>.xlsx
"""
import sys
import os
import openpyxl
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
from openpyxl.utils import get_column_letter

BLACK_FONT  = Font(name="Arial", size=10, color="000000")
BLUE_FONT   = Font(name="Arial", size=10, color="0000FF")
BOLD_FONT   = Font(name="Arial", size=10, bold=True)
TITLE_FONT  = Font(name="Arial", size=14, bold=True)
SUB_FONT    = Font(name="Arial", size=11, bold=True)
GREY_FONT   = Font(name="Arial", size=9, italic=True, color="808080")
HEADER_FONT = Font(name="Arial", size=11, bold=True, color="FFFFFF")
RED_FONT    = Font(name="Arial", size=10, bold=True, color="C00000")

HEADER_FILL   = PatternFill(start_color="000080", end_color="000080", fill_type="solid")
LIGHT_GREEN   = PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid")
LIGHT_YELLOW  = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")
LIGHT_RED     = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")

THIN_BORDER = Border(left=Side(style="thin"), right=Side(style="thin"),
                      top=Side(style="thin"), bottom=Side(style="thin"))
_GRAY_SIDE = Side(style="thin", color="B0B0B0")
INPUT_BORDER = Border(left=_GRAY_SIDE, right=_GRAY_SIDE, top=_GRAY_SIDE, bottom=_GRAY_SIDE)
TOP_BOTTOM = Border(top=Side(style="thin"), bottom=Side(style="double"))
SECTION_BOTTOM = Border(bottom=Side(style="thin"))

FMT_YEN   = '#,##0;(#,##0)'
FMT_PCT   = '0.0%;(0.0%)'
FMT_RATIO = '0.00"x"'
FMT_INT   = '#,##0'


def set_cell(ws, row, col, value, font=None, fmt=None, fill=None, border=None, alignment=None):
    c = ws.cell(row=row, column=col, value=value)
    if font: c.font = font
    if fmt: c.number_format = fmt
    if fill: c.fill = fill
    if border: c.border = border
    if alignment: c.alignment = alignment
    return c


def header_row(ws, row, col_start, labels):
    for i, lbl in enumerate(labels):
        c = ws.cell(row=row, column=col_start + i, value=lbl)
        c.font = HEADER_FONT
        c.fill = HEADER_FILL
        c.alignment = Alignment(horizontal="center", wrap_text=True)
        c.border = SECTION_BOTTOM


def section_title(ws, row, col, text, font=SUB_FONT):
    c = ws.cell(row=row, column=col, value=text)
    c.font = font
    return c


def col_letter(n):
    return get_column_letter(n)


# Research-sourced, as of 2026-07-29/30 (see notes in cells). (name, ticker, market_cap_mn, ownership_pct, status)
LISTED_SUBS = [
    ("AEON Financial Service", "8570.T", 325500, 0.4817, "Listed"),
    ("Tsuruha Holdings",       "3391.T", 1175800, 0.5033, "Listed"),
    ("AEON Kyushu",            "2653.T", 99171,  0.6990, "Listed"),
    ("AEON Hokkaido",          "7512.T", 121016, 0.6562, "Listed"),
    ("Maxvalu Tokai",          "8198.T", 107751, 0.6386, "Listed"),
    ("Ministop",               "9946.T", 52871,  0.4871, "Listed"),
]
DELISTED_SUBS_2025 = [
    ("AEON Mall",   "8905", "Delisted 2025-06-27 (share exchange -> 100% AEON-owned)"),
    ("AEON Delight","9787", "Delisted 2025-07-17 (TOB + squeeze-out -> 100% AEON-owned)"),
    ("Welcia Holdings", "3141", "Delisted 2025-11-27 (merged into Tsuruha Holdings as wholly-owned sub)"),
]


def build_sotp_sheet(wb):
    ws = wb.create_sheet("SOTP Cross-Check")
    ws.sheet_properties.tabColor = "C00000"
    ws.column_dimensions["A"].width = 3
    ws.column_dimensions["B"].width = 30
    for col in "CDEFGH":
        ws.column_dimensions[col].width = 15

    set_cell(ws, 2, 2, "SOTP Cross-Check (intended PRIMARY framework -- see note)", font=TITLE_FONT)
    r = 3
    set_cell(ws, r, 2,
             "AEON is a holding company for many listed/unlisted subsidiaries; a consolidated DCF ('DCF Model' sheet) is an "
             "approximation. SOTP is the theoretically correct lens, but see the critical finding below before using this sheet's totals.",
             font=GREY_FONT)
    r += 2

    set_cell(ws, r, 2, "CRITICAL FINDING: 3 of 9 originally-in-scope subsidiaries went private in 2025", font=BOLD_FONT, fill=LIGHT_RED)
    r += 1
    for name, ticker, note in DELISTED_SUBS_2025:
        set_cell(ws, r, 2, f"{name} ({ticker})", font=BOLD_FONT)
        set_cell(ws, r, 3, note, font=GREY_FONT)
        r += 1
    r += 1
    set_cell(ws, r, 2,
             "Only 6 of the 9 remain listed with a market price. Per the source prompt's own fallback instruction, this sheet "
             "presents the sum of the 6 remaining listed stakes as an explicit FLOOR value -- it does NOT attempt to price the "
             "now-wholly-owned AEON Mall / AEON Delight, the Welcia business now inside Tsuruha, or the GMS/SM retail core "
             "(no reliable, non-guessed EBITDA breakdown by segment is available for AEON). The floor will therefore sit far "
             "below AEON's actual market cap by construction -- that gap IS the (unquantified) value of the wholly-owned businesses.",
             font=GREY_FONT)
    ws.row_dimensions[r].height = 45
    r += 3

    # ── Listed subsidiary stakes ──
    section_title(ws, r, 2, "Listed Subsidiary / Affiliate Stakes")
    r += 1
    header_row(ws, r, 2, ["Company", "Ticker", "Market Cap (JPY mn)", "AEON Ownership %", "Implied Stake Value (JPY mn)"])
    hdr_row = r
    r += 1
    first_data_row = r
    for name, ticker, mcap, pct, status in LISTED_SUBS:
        set_cell(ws, r, 2, name, font=BLACK_FONT)
        set_cell(ws, r, 3, ticker, font=BLACK_FONT)
        set_cell(ws, r, 4, mcap, font=BLUE_FONT, fmt=FMT_YEN, border=INPUT_BORDER)
        set_cell(ws, r, 5, pct, font=BLUE_FONT, fmt=FMT_PCT, border=INPUT_BORDER)
        set_cell(ws, r, 6, f"=D{r}*E{r}", font=BLACK_FONT, fmt=FMT_YEN, border=THIN_BORDER)
        r += 1
    last_data_row = r - 1
    set_cell(ws, r, 2, "Sum of Listed Stakes (JPY mn)", font=BOLD_FONT, fill=LIGHT_GREEN, border=TOP_BOTTOM)
    set_cell(ws, r, 6, f"=SUM(F{first_data_row}:F{last_data_row})", font=BOLD_FONT, fmt=FMT_YEN, fill=LIGHT_GREEN, border=TOP_BOTTOM)
    sum_stakes_row = r
    r += 1
    set_cell(ws, r, 2, "Market cap data as of late Jul 2026; ownership % from FY2026/2 yuho 'related companies' disclosures (Tsuruha stake reflects the Jan-2026 TOB result, 50.11%->50.33% via subsequent market purchases).",
             font=GREY_FONT)
    r += 2

    set_cell(ws, r, 2, "Non-listed / wholly-owned businesses (AEON Mall, AEON Delight, GMS/SM core, Welcia-in-Tsuruha)", font=BOLD_FONT)
    r += 1
    set_cell(ws, r, 2, "NOT PRICED -- no reliable per-segment EBITDA breakdown available; not guessed (per instruction).", font=GREY_FONT, fill=LIGHT_YELLOW)
    r += 2

    # ── Bridge to equity value ──
    section_title(ws, r, 2, "Bridge to SOTP Equity Value (FLOOR -- listed stakes only)")
    r += 1
    set_cell(ws, r, 2, "Sum of Listed Stakes (JPY mn)", font=BOLD_FONT)
    set_cell(ws, r, 3, f"=F{sum_stakes_row}", font=BLACK_FONT, fmt=FMT_YEN)
    bridge_stakes_row = r
    r += 1
    set_cell(ws, r, 2, "Less: Parent (unconsolidated) net interest-bearing debt", font=BOLD_FONT)
    set_cell(ws, r, 3, 3099169, font=BLUE_FONT, fmt=FMT_YEN, border=INPUT_BORDER)
    parent_debt_row = r
    r += 1
    set_cell(ws, r, 2, "  Note: parent-only (単体) net debt was not obtained; using CONSOLIDATED net interest-bearing debt "
                        "(excl. bank deposits, excl. MI addback) as a fallback per prompt instruction -- this OVERSTATES the "
                        "deduction versus a true parent-only figure, making the floor value conservative (too low).",
             font=GREY_FONT)
    ws.row_dimensions[r].height = 30
    r += 1
    set_cell(ws, r, 2, "SOTP Equity Value - FLOOR (JPY mn)", font=BOLD_FONT, fill=LIGHT_GREEN, border=TOP_BOTTOM)
    set_cell(ws, r, 3, f"=C{bridge_stakes_row}-C{parent_debt_row}", font=BOLD_FONT, fmt=FMT_YEN, fill=LIGHT_GREEN, border=TOP_BOTTOM)
    floor_equity_row = r
    r += 1
    set_cell(ws, r, 2, "Shares Outstanding", font=BOLD_FONT)
    set_cell(ws, r, 3, "='DCF Model'!C15", font=BLACK_FONT, fmt=FMT_INT)
    shares_row = r
    r += 1
    set_cell(ws, r, 2, "SOTP Floor Value per Share (JPY, pre-discount)", font=BOLD_FONT)
    set_cell(ws, r, 3, f"=ROUND(C{floor_equity_row}/C{shares_row}*1000000,0)", font=BOLD_FONT, fmt=FMT_YEN)
    floor_per_share_row = r
    r += 2
    set_cell(ws, r, 2,
             "This floor is likely NEGATIVE or small: it deducts group-level debt against only a small slice of group value "
             "(the 6 remaining minority-held listed stakes), deliberately excluding the wholly-owned businesses that generate "
             "most of consolidated EBITDA. A negative or very low floor does NOT mean AEON is worth that little -- it means "
             "this floor construction cannot see most of the company. Treat as a data-availability limitation, not a valuation conclusion.",
             font=GREY_FONT)
    ws.row_dimensions[r].height = 45
    r += 3

    # ── Holding company discount sensitivity ──
    section_title(ws, r, 2, "Holding Company Discount Sensitivity (applied to floor equity value)")
    r += 1
    header_row(ws, r, 2, ["Discount %", "SOTP Equity Value (JPY mn)", "Per Share (JPY)", "vs Current Price (1,424)"])
    disc_hdr_row = r
    r += 1
    discounts = [0.0, 0.10, 0.20, 0.30, 0.40]
    disc_first_row = r
    for d in discounts:
        set_cell(ws, r, 2, d, font=BLUE_FONT, fmt=FMT_PCT, border=INPUT_BORDER)
        set_cell(ws, r, 3, f"=$C${floor_equity_row}*(1-B{r})", font=BLACK_FONT, fmt=FMT_YEN)
        set_cell(ws, r, 4, f"=ROUND(C{r}/$C${shares_row}*1000000,0)", font=BLACK_FONT, fmt=FMT_YEN)
        set_cell(ws, r, 5, f"=D{r}/1424-1", font=BLACK_FONT, fmt=FMT_PCT)
        r += 1
    r += 1
    set_cell(ws, r, 2,
             "No discount level makes the floor value approach the current price of 1,424 -- confirming this floor is not a "
             "usable standalone valuation. It only demonstrates that the value of AEON's remaining minority listed stakes is a "
             "small fraction of the group's total equity value; the balance of AEON's ~3,938,454mn market cap is attributable "
             "to wholly-owned businesses this sheet cannot independently price.",
             font=GREY_FONT)
    ws.row_dimensions[r].height = 45
    r += 3

    # ── MI comparison ──
    section_title(ws, r, 2, "Cross-check: Sum of Listed Stakes vs Non-Controlling Interests (book)")
    r += 1
    set_cell(ws, r, 2, "Sum of Listed Stakes (JPY mn, AEON's share)", font=BOLD_FONT)
    set_cell(ws, r, 3, f"=F{sum_stakes_row}", font=BLACK_FONT, fmt=FMT_YEN)
    mi_compare_row1 = r
    r += 1
    set_cell(ws, r, 2, "Non-Controlling Interests, book value (JPY mn)", font=BOLD_FONT)
    set_cell(ws, r, 3, 984094, font=BLUE_FONT, fmt=FMT_YEN, border=INPUT_BORDER)
    mi_compare_row2 = r
    r += 1
    set_cell(ws, r, 2, "Difference (JPY mn)", font=BOLD_FONT)
    set_cell(ws, r, 3, f"=C{mi_compare_row1}-C{mi_compare_row2}", font=BLACK_FONT, fmt=FMT_YEN)
    diff_row = r
    r += 1
    set_cell(ws, r, 2, "Difference per share (JPY)", font=BOLD_FONT)
    set_cell(ws, r, 3, f"=ROUND(C{diff_row}/C{shares_row}*1000000,0)", font=BLACK_FONT, fmt=FMT_YEN)
    r += 2
    set_cell(ws, r, 2,
             "NOT a clean apples-to-apples comparison: NCI book value (984,094) covers ALL consolidated subsidiaries "
             "(including the now-wholly-owned AEON Mall/AEON Delight, where NCI has since gone to zero post-buyout, and "
             "many unlisted subs), while the listed-stakes sum above covers only the 6 companies still listed today. The "
             "rough numerical proximity between the two figures is not strong evidence that book-value MI is fairly stated; "
             "it is a coincidence of a mixed, shifting subsidiary base, not a controlled comparison.",
             font=GREY_FONT)
    ws.row_dimensions[r].height = 45

    ws.freeze_panes = "C6"
    return {"floor_per_share_row": floor_per_share_row}


def patch_comps_exclude_self(wb):
    if "Comps Analysis" not in wb.sheetnames:
        return False
    ws = wb["Comps Analysis"]
    last_comp_row = 4
    row = 5
    while ws.cell(row=row, column=2).value not in (None, ""):
        last_comp_row = row
        row += 1
    if last_comp_row <= 5:
        return False

    stat_rows = {"25th Percentile": None, "Median (50th)": None, "75th Percentile": None}
    for r in range(1, 40):
        label = ws.cell(row=r, column=2).value
        if label in stat_rows:
            stat_rows[label] = r

    stat_col_map = [(4, 10), (5, 11), (6, 12), (7, 13), (8, 14), (9, 15)]
    patched = 0
    for label, r in stat_rows.items():
        if r is None:
            continue
        for dst_col, src_col in stat_col_map:
            src_letter = col_letter(src_col)
            rng = f"{src_letter}6:{src_letter}{last_comp_row}"
            if label == "25th Percentile":
                formula = f"=PERCENTILE({rng},0.25)"
            elif label == "Median (50th)":
                formula = f"=MEDIAN({rng})"
            else:
                formula = f"=PERCENTILE({rng},0.75)"
            ws.cell(row=r, column=dst_col, value=formula)
            patched += 1
    print(f"  [add_sotp_crosscheck] Patched {patched} Comps Analysis stat formulas to exclude subject row 5 "
          f"(stats now over rows 6:{last_comp_row}).")
    return True


def main():
    if len(sys.argv) < 2:
        print("Usage: python scripts/add_sotp_crosscheck.py <path-to-dcf-model.xlsx>")
        sys.exit(1)
    path = sys.argv[1]
    if not os.path.exists(path):
        print(f"ERROR: file not found: {path}")
        sys.exit(1)

    wb = openpyxl.load_workbook(path)
    if "SOTP Cross-Check" in wb.sheetnames:
        del wb["SOTP Cross-Check"]

    build_sotp_sheet(wb)
    patch_comps_exclude_self(wb)

    wb.save(path)
    print(f"Saved SOTP Cross-Check sheet into: {path}")
    print("Run scripts/recalc_excel_com.py on this file next to compute formula values.")


if __name__ == "__main__":
    main()
