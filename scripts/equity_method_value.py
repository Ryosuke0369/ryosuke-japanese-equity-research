# -*- coding: utf-8 -*-
"""equity_method_value.py — 型F（持分法主導）の「1株あたり別途加算」機構。

なぜ必要か
----------
総合商社は連結利益の相当部分が持分法投資損益である。8058 三菱商事の FY2026/3 は
持分法投資損益 467,941 に対し、売上総利益 − 販管費 で測ったコア営業利益が 418,621 で、
**持分法のほうが大きい**。ところが持分法投資に対応するキャッシュは受取配当だけなので、
連結 FCF を割り引く通常の DCF は投資そのものの価値を取りこぼす。連結 OPM 2.2% で
DCF を組めば EV は構造的に過小になる。

方式（設計は 2026-09-07 のプロンプトで裁定済み）
------------------------------------------------
    1株 Target = コア DCF の 1株値 ＋ 持分法投資価値の 1株あたり

**コア営業利益 = 売上総利益 − 販売費及び一般管理費**。IFRS の商社 P/L では
持分法投資損益・金融収益(受取配当を含む)・投資損益・その他の収益費用は
すべて販管費より下にあるため、この定義だけで

  * 持分法投資損益の除外（加算側との二重計上の防止）
  * 受取配当金の除外（同上。設計書 §2 の要求）

が同時に、かつ銘柄ごとの個別調整なしに満たされる。各社の P/L で持分法損益が
販管費より下にあることは有報で確認して `_core_op_note` に記録すること。

**持分法投資価値**は次の2通り。どちらを採ったかを必ず Adjustments Log に残す。

  `method: "listed_stakes"` … Σ(上場持分先の時価総額 × 議決権比率) ＋ 非上場分簿価。
      非上場分 = 持分法投資の BS 残高 − 上場分の簿価。上場分の簿価が注記から
      取れるときだけ使える。
  `method: "book_value"` …… 持分法投資の BS 残高 × 1.0。上場先の簿価が注記から
      取れない場合はこちらに統一する（設計書の指示）。上場分の時価評価は諦める
      ―― 保守的に出る代わりに、推測が1つも入らない。

二重計上の禁止: 上場分を簿価と時価で二度足さないこと。`listed_stakes` を使う場合、
非上場分は必ず「BS 残高 − 上場分簿価」で求め、BS 残高をそのまま足さない。
本モジュールはこれを構造的に強制する（下の resolve_config）。
"""
import io
import json
import os
import sys

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

BLACK = Font(name="Arial", size=10, color="000000")
BLUE = Font(name="Arial", size=10, color="0000FF")
GREEN = Font(name="Arial", size=10, color="008000")
BOLD = Font(name="Arial", size=10, bold=True)
TITLE = Font(name="Arial", size=14, bold=True)
SUB = Font(name="Arial", size=11, bold=True)
GREY = Font(name="Arial", size=9, italic=True, color="808080")
HDR = Font(name="Arial", size=11, bold=True, color="FFFFFF")

HDR_FILL = PatternFill("solid", start_color="000080", end_color="000080")
GREEN_FILL = PatternFill("solid", start_color="E2EFDA", end_color="E2EFDA")
YELLOW_FILL = PatternFill("solid", start_color="FFF2CC", end_color="FFF2CC")

_T = Side(style="thin")
THIN = Border(left=_T, right=_T, top=_T, bottom=_T)
_G = Side(style="thin", color="B0B0B0")
INPUT = Border(left=_G, right=_G, top=_G, bottom=_G)
TOTAL = Border(top=Side(style="thin"), bottom=Side(style="double"))

YEN = '#,##0;(#,##0)'
PCT = '0.0%;(0.0%)'
RATIO = '0.00"x"'

METHODS = ("book_value", "listed_stakes")


def _cell(ws, r, c, v, font=None, fmt=None, fill=None, border=None):
    x = ws.cell(row=r, column=c, value=v)
    if font:
        x.font = font
    if fmt:
        x.number_format = fmt
    if fill:
        x.fill = fill
    if border:
        x.border = border
    return x


def resolve_config(overrides):
    """`equity_method` ブロックを検証して評価方式を確定する。

    足りない値をデフォルトで埋めることはしない。持分法投資価値は Target の
    構成要素そのものなので、黙って 0 や簿価に落ちると Target が静かに壊れる。
    """
    blk = overrides.get("equity_method")
    if not isinstance(blk, dict):
        raise ValueError(
            "型F: overrides に equity_method ブロックがありません。"
            "持分法投資価値は Target の構成要素であり、既定値に落とせません。")

    bal = blk.get("balance_mn")
    if not isinstance(bal, (int, float)) or isinstance(bal, bool) or bal <= 0:
        raise ValueError(
            "equity_method.balance_mn: 連結BS の『持分法で会計処理されている投資』の"
            "残高（JPY mn、正の数値）が必須です")

    method = str(blk.get("method", "")).strip().lower()
    if method not in METHODS:
        raise ValueError(
            f"equity_method.method: {method!r} は {METHODS} のいずれかでなければ"
            f"なりません（どちらを採ったかは Adjustments Log に残ります）")

    stakes = blk.get("listed_stakes") or []
    listed_book = blk.get("listed_book_mn")
    if method == "listed_stakes":
        if not stakes:
            raise ValueError(
                "equity_method.method='listed_stakes' なのに listed_stakes が空です")
        for s in stakes:
            for k in ("name", "ticker", "market_cap_mn", "ownership"):
                if s.get(k) is None:
                    raise ValueError(f"listed_stakes[{s.get('name')!r}]: {k} が必要です")
        if not isinstance(listed_book, (int, float)) or isinstance(listed_book, bool):
            raise ValueError(
                "equity_method.listed_book_mn が必要です。上場持分先の【簿価】合計で、"
                "非上場分 = balance_mn − listed_book_mn として求めます。"
                "これが注記から取れないなら method を 'book_value' にしてください"
                "（上場分を簿価と時価で二度足さないための構造的な制約です）")
        if listed_book > bal:
            raise ValueError(
                f"listed_book_mn {listed_book:,.0f} が持分法投資残高 {bal:,.0f} を"
                f"超えています。非上場分が負になります")
    elif stakes:
        # 参考として持たせるのは可。ただし評価には使わないことを明示する。
        pass

    return {
        "balance_mn": float(bal),
        "method": method,
        "multiple": float(blk.get("multiple", 1.0)),
        "listed_stakes": stakes,
        "listed_book_mn": float(listed_book) if listed_book is not None else None,
        "label": blk.get("label") or "持分法投資価値",
        "note": blk.get("note") or "",
        "as_of": blk.get("as_of") or "",
    }


def add_sheet(xlsx, cfg, quiet=False):
    """"Equity Method Value" シートを作り、Executive Summary の Target に加算する。"""
    wb = openpyxl.load_workbook(xlsx)
    if "Equity Method Value" in wb.sheetnames:
        del wb["Equity Method Value"]
    ws = wb.create_sheet("Equity Method Value")
    ws.sheet_properties.tabColor = "7030A0"
    ws.column_dimensions["A"].width = 3
    ws.column_dimensions["B"].width = 42
    for c in "CDEFG":
        ws.column_dimensions[c].width = 18

    _cell(ws, 2, 2, "Equity Method Value（型F: 1株あたり別途加算）", font=TITLE)
    r = 3
    _cell(ws, r, 2,
          "連結利益の相当部分が持分法投資損益であり、対応するキャッシュは受取配当のみである。"
          "コア DCF は売上総利益 − 販管費で測ったコア営業利益に基づくため持分法損益を含まず、"
          "その分の価値をこのシートで別途加算する。二重計上を避けるため、"
          "コア側には持分法損益も受取配当も入っていない。", font=GREY)
    ws.row_dimensions[r].height = 30
    r += 2

    stakes_sum_row = None
    if cfg["listed_stakes"]:
        _cell(ws, r, 2, "上場持分先（参考表示）" if cfg["method"] == "book_value"
              else "上場持分先", font=SUB)
        r += 1
        for i, lbl in enumerate(["Company", "Ticker", "Market Cap (JPY mn)",
                                 "議決権比率", "持分時価 (JPY mn)"]):
            c = ws.cell(row=r, column=2 + i, value=lbl)
            c.font, c.fill = HDR, HDR_FILL
            c.alignment = Alignment(horizontal="center", wrap_text=True)
        r += 1
        first = r
        for s in cfg["listed_stakes"]:
            _cell(ws, r, 2, s["name"], font=BLACK, border=THIN)
            _cell(ws, r, 3, s["ticker"], font=BLACK, border=THIN)
            _cell(ws, r, 4, s["market_cap_mn"], font=BLUE, fmt=YEN, border=INPUT)
            _cell(ws, r, 5, s["ownership"], font=BLUE, fmt=PCT, border=INPUT)
            _cell(ws, r, 6, f"=D{r}*E{r}", font=BLACK, fmt=YEN, border=THIN)
            r += 1
        _cell(ws, r, 2, "上場持分の時価合計", font=BOLD, fill=GREEN_FILL, border=TOTAL)
        _cell(ws, r, 6, f"=SUM(F{first}:F{r-1})", font=BOLD, fmt=YEN,
              fill=GREEN_FILL, border=TOTAL)
        stakes_sum_row = r
        r += 2

    _cell(ws, r, 2, "持分法投資価値の算定", font=SUB)
    r += 1

    bal_row = r
    _cell(ws, r, 2, "持分法で会計処理されている投資（連結BS 残高、JPY mn）", font=BLACK, border=THIN)
    _cell(ws, r, 3, cfg["balance_mn"], font=BLUE, fmt=YEN, border=INPUT)
    r += 1

    if cfg["method"] == "book_value":
        mult_row = r
        _cell(ws, r, 2, "評価倍率（簿価に対する倍率）", font=BLACK, border=THIN)
        _cell(ws, r, 3, cfg["multiple"], font=BLUE, fmt=RATIO, border=INPUT)
        r += 1
        val_row = r
        _cell(ws, r, 2, "持分法投資価値（JPY mn）", font=BOLD, fill=GREEN_FILL, border=TOTAL)
        _cell(ws, r, 3, f"=C{bal_row}*C{mult_row}", font=BOLD, fmt=YEN,
              fill=GREEN_FILL, border=TOTAL)
        r += 1
        _cell(ws, r, 2,
              "方式: 簿価（BS残高 × 1.0）。上場持分先の【簿価】が注記から取れないため、"
              "上場分の時価評価は行わない。上場分を簿価と時価で二度足す危険を避ける代わりに、"
              "含み益のある上場持分の分だけ保守的（過小）に出る。", font=GREY, fill=YELLOW_FILL)
        ws.row_dimensions[r].height = 30
        r += 2
    else:
        lb_row = r
        _cell(ws, r, 2, "うち上場持分先の簿価（注記より、JPY mn）", font=BLACK, border=THIN)
        _cell(ws, r, 3, cfg["listed_book_mn"], font=BLUE, fmt=YEN, border=INPUT)
        r += 1
        unl_row = r
        _cell(ws, r, 2, "非上場分の簿価（= BS残高 − 上場分簿価）", font=BLACK, border=THIN)
        _cell(ws, r, 3, f"=C{bal_row}-C{lb_row}", font=BLACK, fmt=YEN, border=THIN)
        r += 1
        mv_row = r
        _cell(ws, r, 2, "上場持分の時価合計", font=BLACK, border=THIN)
        _cell(ws, r, 3, f"=F{stakes_sum_row}", font=GREEN, fmt=YEN, border=THIN)
        r += 1
        val_row = r
        _cell(ws, r, 2, "持分法投資価値（JPY mn）", font=BOLD, fill=GREEN_FILL, border=TOTAL)
        _cell(ws, r, 3, f"=C{mv_row}+C{unl_row}", font=BOLD, fmt=YEN,
              fill=GREEN_FILL, border=TOTAL)
        r += 1
        _cell(ws, r, 2,
              "方式: 上場分は時価（時価総額×議決権比率）、非上場分は簿価。"
              "非上場分は BS残高 − 上場分簿価 で求めており、上場分を簿価と時価で"
              "二度足していない。", font=GREY, fill=YELLOW_FILL)
        ws.row_dimensions[r].height = 22
        r += 2

    sh_row = r
    _cell(ws, r, 2, "希薄化後株式数", font=BLACK, border=THIN)
    _cell(ws, r, 3, "='DCF Model'!C15", font=GREEN, fmt='#,##0', border=THIN)
    r += 1
    ps_row = r
    _cell(ws, r, 2, "1株あたり持分法投資価値（JPY）", font=BOLD, fill=GREEN_FILL, border=TOTAL)
    _cell(ws, r, 3, f"=ROUND(C{val_row}/C{sh_row}*1000000,0)", font=BOLD, fmt=YEN,
          fill=GREEN_FILL, border=TOTAL)
    r += 2

    if cfg["note"]:
        _cell(ws, r, 2, cfg["note"], font=GREY)
        ws.row_dimensions[r].height = 30
        r += 1
    if cfg["as_of"]:
        _cell(ws, r, 2, f"基準日: {cfg['as_of']}", font=GREY)

    ws.freeze_panes = "C4"

    refs = {"value_row": val_row, "per_share_row": ps_row, "balance_row": bal_row}
    wired = wire_exec_summary(wb, refs, cfg, quiet=quiet)
    wb.save(xlsx)
    if not quiet:
        print(f"  [型F] Equity Method Value シートを作成（方式 {cfg['method']}、"
              f"BS残高 {cfg['balance_mn']:,.0f} mn）")
    return {"refs": refs, "exec": wired}


def wire_exec_summary(wb, refs, cfg, quiet=False):
    """Target に「1株あたり持分法投資価値」を加算する。

    行の挿入は Exit 行の【直後】に行う。Target の数式が参照するのは PGM/Exit の行
    （挿入位置より上）なので、openpyxl が数式を書き換えなくても参照は正しいまま
    残る。裁定（追補6 §X）で片脚に降格されている場合も同じで、既存の数式を
    そのまま括弧で包んで加算するため、降格の結果を壊さない。
    """
    if "Executive Summary" not in wb.sheetnames:
        return None
    ws = wb["Executive Summary"]

    def find(prefix, limit=46):
        for r in range(1, limit):
            v = ws.cell(r, 2).value
            if isinstance(v, str) and v.startswith(prefix):
                return r
        return None

    r_tgt = find("Target Price")
    r_pgm = find("DCF - Perpetuity Growth")
    r_exit = find("DCF - Exit Multiple")
    r_note = find("Note: Target Mid")
    if not r_tgt:
        return None

    old = ws.cell(r_tgt, 3).value
    if not isinstance(old, str) or not old.startswith("="):
        raise ValueError(
            f"型F: Executive Summary C{r_tgt} が数式ではありません（{old!r}）。"
            f"コア DCF の Target が数式でないと加算後の値が追跡できません")

    anchor = max([x for x in (r_pgm, r_exit) if x] or [r_tgt])
    ws.insert_rows(anchor + 1, 1)
    r_add = anchor + 1
    if r_note and r_note > anchor:
        r_note += 1
    if r_tgt > anchor:
        r_tgt += 1

    ws.cell(r_add, 2).value = f"{cfg['label']} [1株・別途加算]"
    ws.cell(r_add, 3).value = f"='Equity Method Value'!C{refs['per_share_row']}"
    ws.cell(r_add, 3).number_format = YEN
    ws.cell(r_add, 2).font = BOLD
    ws.cell(r_add, 3).font = BOLD

    # 既存の Target 式（コアDCF）をそのまま包んで加算する。
    ws.cell(r_tgt, 3).value = f"=({old[1:]})+C{r_add}"
    lab = str(ws.cell(r_tgt, 2).value or "Target Price (Mid)")
    if "持分法" not in lab:
        ws.cell(r_tgt, 2).value = lab.split(" (")[0] + " (コアDCF + 持分法投資価値)"

    sentence = (
        "【型F: 持分法主導】コア営業利益 = 売上総利益 − 販売費及び一般管理費 で定義し、"
        "持分法投資損益・受取配当金・金融損益・投資損益をコアから除外している。"
        "除外した持分法投資の価値は上の別途加算行で加算する（二重計上なし）。"
        f"持分法投資価値の算定方式は '{cfg['method']}'、詳細は Equity Method Value シート。"
        "Comps は EV 倍率が持分法混在で歪むため PER/PBR を参照とする。")
    if r_note:
        cur = str(ws.cell(r_note, 2).value or "")
        if "型F" not in cur:
            ws.cell(r_note, 2).value = (cur + " ■" + sentence) if cur else sentence

    if not quiet:
        print(f"  [型F] Executive Summary: Target = (コアDCF) + C{r_add}"
              f"（1株あたり持分法投資価値）")
    return {"target_row": r_tgt, "addon_row": r_add}


def log_to_adjustments(xlsx, cfg, quiet=False):
    """どちらの方式を採ったかを Adjustments Log に残す（設計書の要求）。"""
    wb = openpyxl.load_workbook(xlsx)
    if "Adjustments Log" not in wb.sheetnames:
        return False
    ws = wb["Adjustments Log"]
    r = ws.max_row + 1
    rows = [
        ("equity_method_method", cfg["method"]),
        ("equity_method_balance_mn", cfg["balance_mn"]),
        ("equity_method_multiple", cfg["multiple"]),
        ("equity_method_basis",
         "簿価（BS残高×1.0）。上場先の簿価が注記から取れないため上場分の時価評価は行わない"
         if cfg["method"] == "book_value" else
         f"上場分は時価（{len(cfg['listed_stakes'])}社）、非上場分は BS残高−上場分簿価 "
         f"{cfg['listed_book_mn']:,.0f}"),
    ]
    for k, v in rows:
        ws.cell(r, 2).value = k
        ws.cell(r, 3).value = v
        r += 1
    wb.save(xlsx)
    if not quiet:
        print(f"  [型F] Adjustments Log に評価方式を記録（{cfg['method']}）")
    return True


def main():
    if len(sys.argv) < 2:
        print(__doc__)
        print("Usage: python scripts/equity_method_value.py <xlsx> [--overrides PATH]")
        sys.exit(2)
    xlsx = sys.argv[1]
    ov = None
    if "--overrides" in sys.argv:
        ov = sys.argv[sys.argv.index("--overrides") + 1]
    else:
        base = os.path.basename(xlsx).split("_")[0]
        root = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
        ov = os.path.join(root, "data", "overrides", f"{base}_overrides.json")
    if not os.path.exists(ov):
        print(f"ERROR: overrides not found: {ov}")
        sys.exit(2)
    cfg = resolve_config(json.load(io.open(ov, encoding="utf-8")))
    add_sheet(xlsx, cfg)
    log_to_adjustments(xlsx, cfg)
    print(f"Saved: {xlsx}")
    print("Run scripts/recalc_excel_com.py on this file next.")


if __name__ == "__main__":
    main()
