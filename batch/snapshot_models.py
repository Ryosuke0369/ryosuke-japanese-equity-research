"""Snapshot the decision-relevant cells of a set of DCF workbooks to JSON.

Used to freeze the "before" state of the 2026-09-05 batch before フェーズ2
regenerates every ticker, so the new/old Target comparison (§5) is grounded in
the workbooks themselves rather than in a state file somebody could have edited.
Also used afterwards on the regenerated models, so both sides of the comparison
are read by the same code.

Usage:
    python batch/snapshot_models.py --out batch/phase2_before.json
    python batch/snapshot_models.py --date 20260906 --out batch/phase2_after.json
"""
import argparse
import glob
import json
import os
import sys
import warnings

warnings.filterwarnings("ignore")
sys.stdout.reconfigure(encoding="utf-8", errors="replace")
import openpyxl  # noqa: E402

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)

# (sheet, cell, key). Read from the cached values, so the workbook must have
# been recalculated - an un-recalced one shows up as all-null and is reported.
CELLS = [
    ("Executive Summary", "C9",  "price"),
    ("Executive Summary", "C10", "target_mid"),
    ("Executive Summary", "C11", "verdict"),
    ("Executive Summary", "C12", "upside"),
    ("Executive Summary", "C16", "dcf_pgm"),
    ("Executive Summary", "C17", "dcf_exit"),
    ("Executive Summary", "C18", "comps_ev_ebitda"),
    ("Executive Summary", "C19", "comps_per"),
    ("DCF Model", "C8",  "beta"),
    ("DCF Model", "C10", "size_premium"),
    ("DCF Model", "C13", "terminal_growth"),
    ("DCF Model", "C14", "exit_multiple"),
    ("DCF Model", "C16", "net_debt"),
    ("DCF Model", "C23", "cost_of_equity"),
    ("DCF Model", "C26", "wacc"),
    ("DCF Model", "C53", "ev_pgm"),
    ("DCF Model", "C62", "ev_exit"),
]


def newest_model(code, date=None):
    pat = (f"{code}_DCF_Model_{date}.xlsx" if date
           else f"{code}_DCF_Model_*.xlsx")
    hits = sorted(glob.glob(os.path.join(ROOT, "models", pat)))
    return hits[-1] if hits else None


def snap(path):
    wbv = openpyxl.load_workbook(path, data_only=True)
    wbf = openpyxl.load_workbook(path, data_only=False)
    out = {"file": os.path.basename(path)}
    for sheet, cell, key in CELLS:
        out[key] = wbv[sheet][cell].value if sheet in wbv.sheetnames else None
    # Pipeline Metadata, for provenance (template_rev, beta basis, guidance ...)
    meta = {}
    if "Adjustments Log" in wbf.sheetnames:
        ws, seen = wbf["Adjustments Log"], False
        for row in ws.iter_rows(min_col=2, max_col=3):
            k, v = row[0].value, row[1].value
            if isinstance(k, str) and k.startswith("Pipeline Metadata"):
                seen = True
                continue
            if seen and k not in (None, "key"):
                meta[str(k)] = v
    out["meta"] = {k: meta.get(k) for k in (
        "template_rev", "beta_raw", "beta_blume_adjusted", "beta_adopted_c8",
        "beta_clamped", "guidance_source", "market_data_source",
        "net_debt_source", "capex_pct_basis", "core_ebitda")}
    # Formula-error census on the cached values, so a #DIV/0! is visible here
    errs = {}
    for wsname in wbv.sheetnames:
        for row in wbv[wsname].iter_rows():
            for c in row:
                if isinstance(c.value, str) and c.value.startswith("#"):
                    errs.setdefault(wsname, []).append(f"{c.coordinate}={c.value}")
    out["errors"] = errs
    return out


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--date", default=None,
                    help="only snapshot models with this date stamp")
    ap.add_argument("--state", default=os.path.join(HERE, "batch_state.json"))
    ap.add_argument("--status", default="done")
    ap.add_argument("--out", required=True)
    a = ap.parse_args()

    state = json.load(open(a.state, encoding="utf-8"))
    codes = sorted(k for k, v in state.items()
                   if isinstance(v, dict) and v.get("status") == a.status)

    res, missing, unrecalced, witherr = {}, [], [], []
    for code in codes:
        p = newest_model(code, a.date)
        if not p:
            missing.append(code)
            continue
        try:
            s = snap(p)
        except Exception as e:
            missing.append(f"{code} ({type(e).__name__})")
            continue
        res[code] = s
        if s.get("target_mid") is None and s.get("price") is None:
            unrecalced.append(code)
        if s["errors"]:
            witherr.append(code)

    with open(a.out, "w", encoding="utf-8") as f:
        json.dump(res, f, ensure_ascii=False, indent=1)
    print(f"snapshotted {len(res)}/{len(codes)} -> {a.out}")
    if missing:
        print(f"  no workbook: {', '.join(missing)}")
    if unrecalced:
        print(f"  no cached values (needs recalc): {', '.join(unrecalced)}")
    print(f"  workbooks carrying formula errors: {len(witherr)}"
          + (f" -> {', '.join(witherr)}" if witherr else ""))


if __name__ == "__main__":
    main()
