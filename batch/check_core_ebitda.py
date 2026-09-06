"""Cross-check the workbook's subject EBITDA against the comps CSV self row.

core_ebitda is derived from EDINET (hist_operating_income[-1] + hist_depreciation[-1]).
When EDINET cannot supply depreciation it silently becomes wrong or None, and the
Comps 'Via EV/EBITDA (Median)' reference price goes negative or N/A — with
validate_output still reporting FAIL 0. This check is the guard for that.
"""
import sys, os, csv, glob, warnings
warnings.filterwarnings("ignore")
sys.stdout.reconfigure(encoding="utf-8")
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import openpyxl
from stale_check import assert_fresh, assert_recalculated  # 追補6 §Z / 追補10 §AO

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))

def check(x):
    code = os.path.basename(x)[:4]
    wb = openpyxl.load_workbook(x, data_only=True)
    ws = wb["Comps Analysis"]
    subj_ebitda = impl = None
    for r in range(1, ws.max_row + 1):
        lab = ws.cell(r, 2).value
        if isinstance(lab, str) and lab.startswith("EBITDA (JPY mn)"):
            subj_ebitda = ws.cell(r, 3).value
        if isinstance(lab, str) and lab.startswith("Via EV/EBITDA"):
            impl = ws.cell(r, 3).value
    csv_path = os.path.join(ROOT, "data", "comps", "%s_comps.csv" % code)
    csv_ebitda = None
    if os.path.isfile(csv_path):
        rows = list(csv.DictReader(open(csv_path, encoding="utf-8")))
        v = rows[0]["EBITDA"].strip()
        csv_ebitda = float(v) if v else None
    ok = (subj_ebitda is not None and csv_ebitda is not None
          and abs(subj_ebitda - csv_ebitda) / max(abs(csv_ebitda), 1) < 0.02)
    bad_impl = (impl is None or isinstance(impl, str) or (isinstance(impl, (int, float)) and impl <= 0))
    status = "OK " if (ok and not bad_impl) else "*** MISMATCH ***"
    print(f"{status} {code}: workbook EBITDA={subj_ebitda!r:>14}  csv={csv_ebitda!r:>14}  "
          f"implied EV/EBITDA price={impl!r}")
    return ok and not bad_impl

if __name__ == "__main__":
    paths = sys.argv[1:] or sorted(glob.glob(os.path.join(ROOT, "models", "*_DCF_Model_2026090[56].xlsx")))
    stale = [p for p in paths if not assert_fresh(p)]      # 追補6 §Z
    stale += [p for p in paths if not assert_recalculated(p)]   # 追補10 §AO
    bad = list(dict.fromkeys([p for p in paths if not check(p)] + stale))
    print("\n%d/%d clean" % (len(paths) - len(bad), len(paths)))
