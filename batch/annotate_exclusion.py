"""Append a ticker-specific exclusion sentence to the Executive Summary note.

手順書v2 §6-4 requires the Exec Summary note to say WHICH method entered the
Target average and WHICH was excluded. The template's generic clause explains the
rule but does not name the method for the ticker at hand, so a model whose PGM
leg is INVALID reads as if Target were still a 2-method average.

Generic: the ticker and the sentence come from the command line, never from this
file. Writes plain text only (a value starting with '=' would be stored as a
formula and can corrupt the workbook — lessons.md / 手順書§5-5).

Usage: python batch/annotate_exclusion.py <xlsx> "<sentence>"
"""
import sys, warnings
warnings.filterwarnings("ignore")
sys.stdout.reconfigure(encoding="utf-8")
import openpyxl

def main(path, sentence):
    assert not sentence.startswith("="), "note must not start with '=' (stored as formula)"
    wb = openpyxl.load_workbook(path)
    ws = wb["Executive Summary"]
    row = None
    for r in range(1, ws.max_row + 1):
        v = ws.cell(r, 2).value
        if isinstance(v, str) and v.startswith("Note: Target Mid"):
            row = r
            break
    if row is None:
        raise SystemExit("ERROR: Exec Summary note row not found — refusing to guess a row number.")
    cur = ws.cell(row, 2).value
    if sentence in cur:
        print(f"already annotated (row {row}) — no change")
        return
    ws.cell(row, 2).value = cur + " ■" + sentence
    wb.save(path)
    print(f"annotated Executive Summary B{row}:\n  {ws.cell(row, 2).value}")
    print("NOTE: cached values were dropped on save — run recalc_excel_com.py then validate_output.py.")

if __name__ == "__main__":
    main(sys.argv[1], sys.argv[2])
