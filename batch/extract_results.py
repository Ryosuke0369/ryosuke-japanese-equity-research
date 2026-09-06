"""Pull the Executive Summary / WACC / scenario numbers out of a generated DCF."""
import sys, os, warnings, glob
warnings.filterwarnings("ignore")
sys.stdout.reconfigure(encoding="utf-8")
import openpyxl

def dump(path):
    wv = openpyxl.load_workbook(path, data_only=True)
    wf = openpyxl.load_workbook(path, data_only=False)
    print(f"### {os.path.basename(path)}")
    es = wv["Executive Summary"]; esf = wf["Executive Summary"]
    for r in range(1, es.max_row + 1):
        lab = es.cell(r, 2).value or es.cell(r, 1).value
        val = es.cell(r, 3).value
        if lab and val is not None:
            print(f"  ES B{r}: {str(lab)[:60]:<62} = {val}")
    d = wv["DCF Model"]; df = wf["DCF Model"]
    print("  -- WACC block --")
    for r in range(1, 40):
        lab = d.cell(r, 2).value
        if lab and isinstance(lab, str) and any(k in lab for k in
            ("Risk", "Beta", "Equity Risk", "Size", "Cost of Debt", "Tax", "D/E",
             "WACC", "Cost of Equity", "Terminal", "Exit", "Capex", "D&A", "Net Debt",
             "Shares", "LTM")):
            print(f"    C{r} {lab[:44]:<46} = {d.cell(r,3).value}")
    print("  -- scenario table --")
    for r in range(60, 90):
        b = d.cell(r, 2).value
        if b in ("Base", "Upside", "Management", "Downside 1", "Downside 2"):
            vals = [d.cell(r, c).value for c in range(3, 10)]
            print(f"    {b:<12} {vals}")
    if "Reverse DCF" in wv.sheetnames:
        rv = wv["Reverse DCF"]
        print("  -- Reverse DCF Block F --")
        for r in range(1, rv.max_row + 1):
            for c in range(1, 6):
                v = rv.cell(r, c).value
                if isinstance(v, str) and ("一行" in v or "Block F" in v or
                                           "one-line" in v.lower() or "answer" in v.lower()):
                    for rr in range(r, min(r + 6, rv.max_row + 1)):
                        row = [rv.cell(rr, cc).value for cc in range(1, 7)]
                        row = [x for x in row if x is not None]
                        if row: print(f"    {row}")
                    return

if __name__ == "__main__":
    for p in sys.argv[1:]:
        dump(p)
